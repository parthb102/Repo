#!/usr/bin/env python3
"""Rebuild slide 2, the framing page: two track cards with aligned rows
   (Objective / Approach / Key assets / Outputs), each Approach bullet labelled
   with the exact phase name and timing used on the workplan pages."""
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import vgkit as k
from pptx.dml.color import RGBColor

F = "working_vg.pptx"
C = json.load(open("content.json"))
FR = C["framing"]

CARD_W = 5.60
CX = [1.35, 7.13]        # card x positions
LBL_X = 0.61               # row label gutter
PAD = 0.22                 # inner padding of a card


def measure(rows_text, w, size, bold_len=0):
    return k.est_lines(rows_text, w, size, bold_prefix_len=bold_len)


def main():
    prs = Presentation(F)
    s = list(prs.slides)[1]

    k.wipe(s, keep_names=("Object 6", "2. Slide Title", "Subtitle 140",
                          "1. On-page tracker", "think-cell data - do not delete"))
    k.set_title(s, FR["title"], FR["subtitle"])
    for sh in s.shapes:
        if sh.name == "2. Slide Title":
            sh.left, sh.top = Inches(0.61), Inches(0.16)
            sh.width, sh.height = Inches(9.70), Inches(0.74)
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(20)
        if sh.name == "Subtitle 140":
            sh.left, sh.top = Inches(0.61), Inches(0.94)
            sh.width, sh.height = Inches(9.70), Inches(0.26)
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(11)

    tw = CARD_W - 2 * PAD          # usable text width inside a card

    # ---------- measure every row so the two cards stay aligned ----------
    H = {}
    obj_h = 0
    for col in FR["columns"]:
        n = measure(col["objective_lead"] + col["objective_rest"], tw, 9,
                    bold_len=len(col["objective_lead"]))
        obj_h = max(obj_h, n * 0.148)
    H["obj"] = obj_h + 0.06

    for key, size, gap in (("approach", 8, 0.06), ("assets", 8, 0.055), ("outputs", 8, 0.055)):
        tot = 0
        for col in FR["columns"]:
            h = 0
            for item in col[key]:
                if isinstance(item, list):
                    txt, bl = item[0] + item[1], len(item[0])
                else:
                    txt, bl = item, 0
                h += measure(txt, tw - 0.10, size, bold_len=bl) * 0.130 + gap
            tot = max(tot, h)
        H[key] = tot + 0.04

    # ---------- geometry ----------
    HEAD_Y, HEAD_H = 1.40, 0.50
    y0 = HEAD_Y + HEAD_H + 0.14
    rows = [("Objective", "obj"), ("Approach", "approach"),
            ("Key assets", "assets"), ("Outputs", "outputs")]
    ROWGAP = 0.13
    body_h = sum(H[k2] for _, k2 in rows) + ROWGAP * (len(rows) - 1) + 0.24
    card_bottom = y0 + body_h

    # ---------- cards ----------
    for ci, col in enumerate(FR["columns"]):
        x = CX[ci]
        k.rect(s, x, y0 - 0.10, CARD_W, body_h + 0.10, fill=RGBColor(0xF7, 0xFA, 0xFE))
        head = k.rect(s, x, HEAD_Y, CARD_W, HEAD_H, fill=(k.CYAN if ci == 0 else k.MED))
        tf = head.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        k.put(tf.paragraphs[0], [(col["header"], 13, True, k.WHITE, False)],
              align=PP_ALIGN.CENTER)
        k.rect(s, x, HEAD_Y + HEAD_H, CARD_W, 0.055, fill=k.NAVY)

    # ---------- row labels + separators + content ----------
    y = y0 + 0.06
    for ri, (label, key) in enumerate(rows):
        lb = k.tbox(s, LBL_X, y, 0.78, 0.24)
        k.put(lb.text_frame.paragraphs[0], [(label, 9.5, True, k.NAVY, False)], ls=0.95)
        if ri:
            for ci in range(2):
                ln = s.shapes.add_connector(
                    1, Inches(CX[ci] + PAD), Inches(y - ROWGAP / 2 - 0.02),
                    Inches(CX[ci] + CARD_W - PAD), Inches(y - ROWGAP / 2 - 0.02))
                ln.line.color.rgb = k.BORD
                ln.line.width = Pt(0.75)
                ln.shadow.inherit = False
            ln = s.shapes.add_connector(1, Inches(LBL_X), Inches(y - ROWGAP / 2 - 0.02),
                                        Inches(LBL_X + 0.78), Inches(y - ROWGAP / 2 - 0.02))
            ln.line.color.rgb = k.BORD
            ln.line.width = Pt(0.75)
            ln.shadow.inherit = False

        for ci, col in enumerate(FR["columns"]):
            x = CX[ci] + PAD
            if key == "obj":
                tb = k.tbox(s, x, y, tw, H["obj"])
                k.put(tb.text_frame.paragraphs[0],
                      [(col["objective_lead"], 9, True, k.NAVY, False),
                       (col["objective_rest"], 9, False, k.BLACK, False)], ls=1.0)
            else:
                yy = y
                gap = 0.06 if key == "approach" else 0.055
                for item in col[key]:
                    if isinstance(item, list):
                        lead, rest = item
                    else:
                        lead, rest = "", item
                    n = measure(f"{lead}{rest}", tw - 0.10, 8, bold_len=len(lead))
                    h = n * 0.130 + 0.028
                    b = k.tbox(s, x, yy, tw, h)
                    p = b.text_frame.paragraphs[0]
                    p.line_spacing = 0.95
                    pPr = p._p.get_or_add_pPr()
                    pPr.set('marL', '82296'); pPr.set('indent', '-82296')
                    runs = [("• ", 8, False, k.BLACK, False)]
                    if lead:
                        runs.append((lead, 8, True, k.NAVY, False))
                    runs.append((rest, 8, False, k.BLACK, False))
                    k.put(p, runs, ls=0.95)
                    yy += h + gap
        y += H[key] + ROWGAP

    # ---------- closing band ----------
    BY = card_bottom + 0.16
    band = k.rect(s, LBL_X, BY, 12.12, 0.44, fill=k.NAVY, rounded=True, adj=0.20)
    tf = band.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    k.put(tf.paragraphs[0], [(FR["closingLine"], 9.5, False, k.WHITE, False)],
          ls=0.95, align=PP_ALIGN.CENTER)

    k.sticky(s, C["stickies"]["framing"], h=1.15)
    prs.save(F)
    print(f"framing built; card bottom {card_bottom:.2f}, band to {BY + 0.44:.2f}")
    print("row heights:", {a: round(b, 2) for a, b in H.items()})


if __name__ == "__main__":
    main()
