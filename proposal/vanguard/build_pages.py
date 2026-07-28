#!/usr/bin/env python3
"""Build the two Vanguard IDP workplan pages (slides 3 and 4) in the Cetera-p30
   grammar: proportional numbered phase chevrons on top, matching numbered
   activity columns beneath, deliverables strip, governance rail."""
import json, sys, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import vgkit as k

F = "working_vg.pptx"
CONTENT = json.load(open("content.json"))


def build_workplan(s, spec, weights, sticky_text):
    """spec: {title, subtitle, phases:[{number,name,timing,activities:[{lead,detail}]}],
              deliverables:[...], footer}"""
    k.wipe(s, keep_names=("Object 2", "Object 6", "2. Slide Title", "Title 2",
                          "1. On-page tracker", "3. Subtitle", "think-cell data - do not delete"))
    k.set_title(s, spec["title"], None)
    # keep the title clear of the review sticky and strip any inherited bullet glyph
    for sh in s.shapes:
        if sh.name in ("2. Slide Title", "Title 2"):
            sh.width = Inches(9.70)
            sh.top = Inches(0.22)
            sh.height = Inches(0.82)
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(20)
                pPr = p._p.get_or_add_pPr()
                for tag in ("a:buChar", "a:buAutoNum"):
                    for el in pPr.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag.split(':')[1]}"):
                        pPr.remove(el)
                pPr.append(pPr.makeelement(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}buNone", {}))
    # these slides have no subtitle placeholder, so add one under the title
    sub = k.tbox(s, k.L, 1.10, 9.70, 0.26)
    k.put(sub.text_frame.paragraphs[0], [(spec["subtitle"], 10, False, k.MED, False)], ls=1.0)

    GUT = 0.98                                   # left gutter for row labels
    cols = k.col_geometry(weights, gap=0.10, left=k.L + GUT, total=k.W - GUT)

    # ---------- phase chevrons ----------
    CH_Y, CH_H = 1.46, 0.54
    OVER = 0.13                      # chevrons interlock
    for i, (ph, (cx, cw)) in enumerate(zip(spec["phases"], cols)):
        fill = k.RAMP[i] if i < len(k.RAMP) else k.NAVY
        x = cx - (0 if i == 0 else OVER)
        w = cw + (OVER if i == 0 else OVER + 0.10)
        ch = k.chevron(s, x, CH_Y, w, CH_H, fill, first=(i == 0))
        tf = ch.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.20 if i else 0.10)
        tf.margin_right = Inches(0.14)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        k.put(tf.paragraphs[0], [(ph["name"], 9.5, True, k.WHITE, False)],
              ls=0.9, align=PP_ALIGN.CENTER)
        k.put(tf.add_paragraph(), [(ph["timing"], 8, False, k.WHITE, False)],
              ls=0.9, align=PP_ALIGN.CENTER)

    # ---------- activity columns (numbered badge ties each to its chevron) ----------
    TOP = CH_Y + CH_H + 0.40
    lab = k.tbox(s, k.L, TOP - 0.10, GUT - 0.06, 0.4)
    k.put(lab.text_frame.paragraphs[0], [("Activities", 9, True, k.GREY, False)], ls=0.95)

    BOT = 0.0
    for i, (ph, (cx, cw)) in enumerate(zip(spec["phases"], cols)):
        fill = k.RAMP[i] if i < len(k.RAMP) else k.NAVY
        k.badge(s, cx + 0.13, TOP - 0.12, ph["number"], fill)
        y = TOP + 0.10
        for act in ph["activities"]:
            _, h = k.bullet(s, cx, y, cw, act["lead"], act.get("detail", ""),
                            size=9, lead_col=k.NAVY, col=k.BLACK)
            y += h + 0.095
        BOT = max(BOT, y)

    # ---------- deliverables strip ----------
    DY = BOT + 0.26
    heights = []
    for (cx, cw), d in zip(cols, spec["deliverables"]):
        heights.append(k.est_lines(d, cw - 0.26, 8.5) * 0.145 + 0.22)
    DH = max(heights)
    k.rect(s, k.L, DY, k.W, DH, fill=k.TINT)
    dl = k.tbox(s, k.L + 0.08, DY + (DH - 0.22) / 2, GUT - 0.10, 0.24)
    k.put(dl.text_frame.paragraphs[0], [("Deliverables", 9, True, k.NAVY, False)])
    for i, ((cx, cw), d) in enumerate(zip(cols, spec["deliverables"])):
        fill = k.RAMP[i] if i < len(k.RAMP) else k.NAVY
        k.rect(s, cx - 0.10, DY + 0.07, 0.045, DH - 0.14, fill=fill)
        tb = k.tbox(s, cx, DY + 0.11, cw, DH - 0.20)
        k.put(tb.text_frame.paragraphs[0], [(d, 8.5, False, k.BLACK, False)], ls=0.95)

    # ---------- governance rail ----------
    GY = DY + DH + 0.16
    ln = s.shapes.add_connector(1, Inches(k.L), Inches(GY), Inches(k.R), Inches(GY))
    ln.line.color.rgb = k.BORD
    ln.line.width = Pt(0.75)
    ln.shadow.inherit = False
    fb = k.tbox(s, k.L, GY + 0.07, k.W, 0.24)
    k.put(fb.text_frame.paragraphs[0], [(spec["footer"], 8, False, k.GREY, True)], ls=0.95)

    k.sticky(s, sticky_text, h=1.05)
    return DY + DH


def main():
    prs = Presentation(F)
    slides = list(prs.slides)

    end3 = build_workplan(
        slides[2], CONTENT["track1"], weights=[2.30, 2.85, 3.95, 2.75],
        sticky_text=CONTENT["stickies"]["track1"])
    print("slide 3 content bottom:", round(end3, 2))

    end4 = build_workplan(
        slides[3], CONTENT["track2"], weights=[2.10, 2.42, 2.42, 2.42, 2.42],
        sticky_text=CONTENT["stickies"]["track2"])
    print("slide 4 content bottom:", round(end4, 2))

    prs.save(F)
    print("saved", F)


if __name__ == "__main__":
    main()
