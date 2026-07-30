#!/usr/bin/env python3
"""Apply Megha's feedback and consistency fixes to the 07/30 EP deck:
   1. Asset 5 swap: Celonis process intelligence + AI sizing -> AA tool and
      AI process workflow tool (slides 7, 8, 9, 13), drop the 'Not confirmed'
      stickies that requested exactly this change.
   2. Pages 4/5: rename 'How we analyze it' -> 'Key analysis', add an
      'Assets deployed' column; retitle both pages so (1/2)/(2/2) describe
      their actual content.
   3. Slide 3: fold 'Spans and layers' into F. Org structure and re-letter
      so the ladder matches the deep-dive pages (A-K, same letters).
   4. Fact fixes from the source-doc check are applied via FACTS dict.
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

F = "work_ep.pptx"
NAVY = RGBColor(0x05, 0x1C, 0x2C)
BLUE = RGBColor(0x14, 0x3C, 0xCC)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x63, 0x66, 0x6A)
YELL = RGBColor(0xFF, 0xFF, 0x66)


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


def find_text(slide, needle):
    for sh in walk(slide.shapes):
        if sh.has_text_frame and needle in sh.text_frame.text:
            yield sh


def set_text(sh, new, size=None, bold=None, color=None, italic=None):
    """Replace a shape's text, inheriting the first run's full rPr (size, bold,
    theme color, font) so replacements keep the template styling."""
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    rpr = copy.deepcopy(p0.runs[0]._r.get_or_add_rPr()) if p0.runs else None
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    first = True
    for ln in new.split("\n"):
        p = p0 if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = ln
        if rpr is not None:
            r._r.insert(0, copy.deepcopy(rpr))
        f = r.font
        if size is not None:
            f.size = Pt(size)
        if bold is not None:
            f.bold = bold
        if color is not None:
            f.color.rgb = color
        if italic is not None:
            f.italic = italic


def del_shape(sh):
    sh._element.getparent().remove(sh._element)


def tbox(slide, x, y, w, h):
    return slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))


def put(par, runs, ls=1.0):
    par.line_spacing = ls
    for txt, sz, bold, col in runs:
        r = par.add_run()
        r.text = txt
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.color.rgb = col
        r.font.name = "Arial"


def sticky(slide, text, x=10.55, y=0.10, w=2.62, h=1.00):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = YELL
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    put(tf.paragraphs[0], [(text, 8, False, BLACK)], ls=0.95)
    return s


# ---- facts corrected after the source-document check (id: old -> new) ----
import json
FACTS = json.load(open("fact_fixes.json"))


def apply_text_swaps(prs, swaps):
    """swaps: list of (slide_idx0, old, new) applied to any shape containing old."""
    hits = 0
    for si, old, new in swaps:
        s = prs.slides[si]
        for sh in walk(s.shapes):
            if sh.has_text_frame and old in sh.text_frame.text:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if old in r.text:
                            r.text = r.text.replace(old, new)
                            hits += 1
                # cross-run occurrences: rebuild if still present
                if old in sh.text_frame.text:
                    set_text(sh, sh.text_frame.text.replace(old, new))
                    hits += 1
    return hits


def main():
    prs = Presentation(F)
    S = prs.slides

    # ---------------- 1. slide 3: re-letter and fold spans into F ------------
    s3 = S[2]
    for sh in walk(s3.shapes):
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t == "G. Spans and layers\n/ gearing ratio" or t.startswith("G. Spans and layers"):
            set_text(sh, "G. Location mix")
        elif t == "H. Location mix":
            del_shape(sh)
        elif t == "F. Org structure":
            set_text(sh, "F. Org structure, spans and layers")
        elif t == "I. FinOps":
            set_text(sh, "H. FinOps")
        elif t == "J. Supplier management":
            set_text(sh, "I. Supplier management")
        elif t == "K. Demand policies":
            set_text(sh, "J. Demand policies")
        elif t == "L. Process streamlining":
            set_text(sh, "K. Process streamlining")

    # savings ranges per fact check (keyed by shape name: two shapes share text)
    # mirror the untouched range shapes' styling (blue bold) on the replacements
    ref_font = None
    for sh in walk(s3.shapes):
        if sh.has_text_frame and sh.name == "Text 13" and sh.text_frame.paragraphs[0].runs:
            ref_font = sh.text_frame.paragraphs[0].runs[0].font
    for sh in walk(s3.shapes):
        if sh.has_text_frame and sh.name in FACTS.get("s3_ranges_by_shape", {}):
            set_text(sh, FACTS["s3_ranges_by_shape"][sh.name])
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if ref_font is not None:
                        if ref_font.size: r.font.size = ref_font.size
                        r.font.bold = ref_font.bold
                        from pptx.enum.dml import MSO_THEME_COLOR
                        if ref_font.color.type is not None and str(ref_font.color.type) == "SCHEME (2)":
                            r.font.color.theme_color = ref_font.color.theme_color
                        else:
                            try:
                                r.font.color.rgb = ref_font.color.rgb
                            except Exception:
                                pass
    if FACTS.get("s3_note"):
        for sh in find_text(s3, "Note: Each lever"):
            set_text(sh, FACTS["s3_note"])
            sh.height = Inches(0.55)
            sh.text_frame.word_wrap = True
    # remove the now-empty chip background where 'H. Location mix' sat, matching
    # by absolute position computed through group transforms
    def walk_abs(shapes, gx=0.0, gy=0.0, sx=1.0, sy=1.0):
        a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for sh in shapes:
            if sh.shape_type == 6:
                xf = sh._element.grpSpPr.find(a + "xfrm")
                ch = xf.find(a + "chOff"); ext = xf.find(a + "ext"); che = xf.find(a + "chExt")
                fx = int(ext.get("cx")) / int(che.get("cx"))
                fy = int(ext.get("cy")) / int(che.get("cy"))
                yield from walk_abs(sh.shapes, sh.left - int(ch.get("x")) * fx,
                                    sh.top - int(ch.get("y")) * fy, fx, fy)
            else:
                yield sh, gx + (sh.left or 0) * sx, gy + (sh.top or 0) * sy
    for sh, ax, ay in list(walk_abs(s3.shapes)):
        if (sh.has_text_frame and not sh.text_frame.text.strip()
                and abs(ax / 914400 - 5.54) < 0.15 and abs(ay / 914400 - 4.38) < 0.15):
            del_shape(sh)
    sticky(s3, "Ranges re-checked vs sources: workforce productivity and operating model were "
               "swapped vs the Rapid IT diagnostic lever table ('Operating Model 10-15%' = ways "
               "of working; 'Talent & Org 10-25%' = org shape and location). Corrected. Tech IDP "
               "quotes 20-30% for full engineering excellence programs; automation's 10-25% is "
               "the thinnest-supported range.", x=7.30, y=0.06, w=2.90, h=1.38)

    # ---------------- 2. pages 4/5 rebuild with 5 columns -------------------
    COLS = [("Lever", 0.56, 1.62), ("What we measure", 2.28, 2.30),
            ("Key analysis", 4.68, 2.50), ("Assets deployed", 7.28, 1.86),
            ("Illustrative outcome", 9.24, 2.52)]

    def rebuild_lever_page(s, title, rows, note):
        # wipe everything except the title placeholder (old grids differ per
        # page: grouped on p4, loose top-level textboxes and rules on p5)
        for sh in list(s.shapes):
            if "Title" not in sh.name:
                del_shape(sh)
        for sh in walk(s.shapes):
            if sh.has_text_frame and "Title" in sh.name:
                set_text(sh, title, size=18, bold=True)
                sh.left, sh.top = Inches(0.56), Inches(0.26)
                sh.width, sh.height = Inches(11.9), Inches(0.55)
        # header row
        HY = 1.06
        for name, x, w in COLS:
            hb = tbox(s, x, HY, w, 0.30)
            put(hb.text_frame.paragraphs[0], [(name, 12, True, BLACK)])
            hb.text_frame.word_wrap = True
        ln = s.shapes.add_connector(1, Inches(0.56), Inches(HY + 0.34), Inches(11.80), Inches(HY + 0.34))
        ln.line.color.rgb = BLACK
        ln.line.width = Pt(1.4)
        ln.shadow.inherit = False
        # body rows
        y = HY + 0.46
        n = len(rows)
        rh = (6.55 - y) / n
        for ri, row in enumerate(rows):
            for (name, x, w), txt in zip(COLS, row):
                b = tbox(s, x, y + 0.02, w, rh - 0.10)
                b.text_frame.word_wrap = True
                first = True
                for para in txt.split("\n"):
                    p = b.text_frame.paragraphs[0] if first else b.text_frame.add_paragraph()
                    first = False
                    bold = (name == "Lever")
                    put(p, [(para, 10.5, bold, BLACK)], ls=1.0)
            if ri:
                ln = s.shapes.add_connector(1, Inches(0.56), Inches(y - 0.06), Inches(11.80), Inches(y - 0.06))
                ln.line.color.rgb = RGBColor(0xBF, 0xBF, 0xBF)
                ln.line.width = Pt(0.75)
                ln.shadow.inherit = False
            y += rh
        if note:
            sticky(s, note)

    R4 = FACTS["page4_rows"]
    R5 = FACTS["page5_rows"]
    rebuild_lever_page(S[3], "Lever deep dive (1/2): demand management and workforce productivity", R4,
                       "Added per Megha: Key analysis and Assets deployed columns. Appendix of sample outputs per asset still to come - Vamsi.")
    rebuild_lever_page(S[4], "Lever deep dive (2/2): operating model, expense management and automation", R5, None)

    # ---------------- 3. asset 5 swap on slides 7, 8, 9, 13 ------------------
    A5 = "AA tool and AI process workflow tool"
    s7, s8, s9, s13 = S[6], S[7], S[8], S[12]

    def match_run_style(slide, src_name, dst_sh):
        """Give dst_sh's runs the same rPr as slide's src_name first run (the
        old asset-5 lines were red-flagged 'not confirmed'; restyle to match
        their confirmed siblings)."""
        src = next((x for x in walk(slide.shapes) if x.name == src_name and x.has_text_frame
                    and x.text_frame.paragraphs[0].runs), None)
        if src is None:
            return
        rpr = src.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
        for p in dst_sh.text_frame.paragraphs:
            for r in p.runs:
                old = r._r.find(qn("a:rPr"))
                if old is not None:
                    r._r.remove(old)
                r._r.insert(0, copy.deepcopy(rpr))

    for sh in find_text(s7, "Celonis process intelligence and AI sizing"):
        set_text(sh, A5)
        match_run_style(s7, "Text 7", sh)
    for sh in find_text(s7, "Process baselines and the AI value at stake by business domain"):
        set_text(sh, "AI value at stake by domain, then re-imagined workflows for the priority domains")
    for sh in find_text(s7, "Process mining plus structured value sizing"):
        set_text(sh, "Internal value sizing, then AI-assisted workflow redesign")
    for sh in find_text(s8, "Celonis process intelligence and AI sizing"):
        set_text(sh, A5)
    for sh in find_text(s8, "Process baselines and value at stake"):
        set_text(sh, "Value at stake and re-imagined workflows")
    for sh in find_text(s9, "5  Celonis process intelligence and AI sizing"):
        set_text(sh, "5  " + A5)
        match_run_style(s9, "TextBox 7", sh)
    for sh in find_text(s9, "Current-state process baselines and the AI value at stake by business domain"):
        set_text(sh, "AI value at stake by business domain, and re-imagined future-state workflows for the priority domains")
    for sh in find_text(s9, "Process mining on system logs, then structured value sizing by domain"):
        set_text(sh, "AA tool sizing on domain data (run internally), then AI process workflow tool for redesign")
    for sh in find_text(s9, "McKinsey AI value-at-stake sizing tool"):
        set_text(sh, "McKinsey value-at-stake benchmarks")
    # table on s13
    for sh in walk(s13.shapes):
        if getattr(sh, 'has_table', False) and sh.has_table:
            tbl = sh.table
            for r in tbl.rows:
                c0 = r.cells[0]
                if "Celonis" in c0.text:
                    for p in list(c0.text_frame.paragraphs[1:]):
                        p._p.getparent().remove(p._p)
                    for run in list(c0.text_frame.paragraphs[0].runs):
                        run._r.getparent().remove(run._r)
                    rr = c0.text_frame.paragraphs[0].add_run()
                    rr.text = "5  " + A5
                    rr.font.size = Pt(11)
                    rr.font.bold = True
                    c1 = r.cells[1]
                    for p in list(c1.text_frame.paragraphs[1:]):
                        p._p.getparent().remove(p._p)
                    for run in list(c1.text_frame.paragraphs[0].runs):
                        run._r.getparent().remove(run._r)
                    rr = c1.text_frame.paragraphs[0].add_run()
                    rr.text = ("AI value at stake by domain from the AA tool  |  Re-imagined future-state "
                               "workflows for the priority domains from the AI process workflow tool  |  "
                               "A value and feasibility heatmap to choose the domains to reimagine first")
                    rr.font.size = Pt(11)
    # drop the two 'Not confirmed' stickies now that the swap is applied
    for s in (s7, s9):
        for sh in list(s.shapes):
            if sh.name == "Sticky" and "Not confirmed" in (sh.text_frame.text if sh.has_text_frame else ""):
                del_shape(sh)

    # ---------------- 4. remaining fact fixes (benchmark stats) --------------
    n = apply_text_swaps(prs, [(t["slide"] - 1, t["old"], t["new"]) for t in FACTS.get("swaps", [])])
    print("fact swaps applied:", n)

    # s12 caption re-attribution, keyed by shape name (duplicate caption text on page)
    s12 = S[11]
    for sh in walk(s12.shapes):
        if sh.has_text_frame and sh.name in FACTS.get("s12_caption_by_shape", {}):
            set_text(sh, FACTS["s12_caption_by_shape"][sh.name])

    # swap the duplicate AIQ tile (Image 4) for the Core Tech IDP lever value-sizing
    # exhibit, and re-caption it; no DVI or AA-tool sample exists in the source docs
    from pptx.util import Emu
    for sh in list(walk(s12.shapes)):
        if sh.shape_type == 13 and sh.name == "Image 4":
            L, T, W, H = sh.left, sh.top, sh.width, sh.height
            sh._element.getparent().remove(sh._element)
            s12.shapes.add_picture("swap_valuesizing.png", L, T, W, H)
    cap_font = None
    for sh in walk(s12.shapes):
        if sh.has_text_frame and sh.name == "Text 5" and sh.text_frame.paragraphs[0].runs:
            cap_font = sh.text_frame.paragraphs[0].runs[0].font
    def style_caption(sh):
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if cap_font is None:
                    continue
                if cap_font.size: r.font.size = cap_font.size
                r.font.bold = cap_font.bold
                if cap_font.color.type is not None and str(cap_font.color.type) == "SCHEME (2)":
                    r.font.color.theme_color = cap_font.color.theme_color
                else:
                    try:
                        r.font.color.rgb = cap_font.color.rgb
                    except Exception:
                        pass
    for sh in walk(s12.shapes):
        if not sh.has_text_frame:
            continue
        if sh.name == "Text 17":
            set_text(sh, "Value sizing")
            style_caption(sh)
        elif sh.name == "Text 18":
            set_text(sh, "Opportunity sized by lever, with difficulty and timeframe")
        elif sh.name in FACTS.get("s12_caption_by_shape", {}):
            style_caption(sh)
    sticky(s12, "Screenshot provenance verified: all six tiles come from the shared docs "
                "(Rapid IT diagnostic p14/17/19, AIQ p24, OMI p10, Core Tech IDP p9). "
                "No DVI or AA tool sample outputs exist in any doc we were given, so the "
                "duplicate AIQ tile became the lever value-sizing exhibit. Ask Jen for a DVI "
                "sample page and an AI value-at-stake sample to complete asset coverage.",
           x=10.45, y=0.06, w=2.76, h=1.30)

    # cover: this doc is the levers/assets/outputs discussion, not the planning
    # doc; fix subtitle (also removes the en dash) and refresh the date
    for sh in walk(S[0].shapes):
        if sh.has_text_frame and "Planning document" in sh.text_frame.text:
            set_text(sh, "Technology productivity: levers, assets to deploy, and expected outputs\nJuly 30, 2026")

    prs.save(F)
    print("saved", F)


if __name__ == "__main__":
    main()
