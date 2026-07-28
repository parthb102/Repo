#!/usr/bin/env python3
"""Shared layout kit for the Vanguard IDP deck.
   Cetera-p30 grammar: proportional phase chevrons with number badges on top,
   matching numbered activity columns beneath, deliverables strip, governance rail.
   McKinsey theme colors pulled from the deck's own theme1.xml."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# --- deck theme (theme1.xml) ---
NAVY   = RGBColor(0x06, 0x1F, 0x79)   # accent1
CYAN   = RGBColor(0x00, 0xA9, 0xF4)   # accent2
BLUE   = RGBColor(0x22, 0x51, 0xFF)   # accent3
PALE   = RGBColor(0x99, 0xE6, 0xFF)   # accent4
MED    = RGBColor(0x06, 0x79, 0xC3)   # accent5
GREYC  = RGBColor(0x75, 0x78, 0x7B)
GREY   = RGBColor(0x63, 0x66, 0x6A)
LGREY  = RGBColor(0xF2, 0xF2, 0xF2)
TINT   = RGBColor(0xEF, 0xF6, 0xFE)
BORD   = RGBColor(0xD9, 0xD9, 0xD9)
BLACK  = RGBColor(0, 0, 0)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
YELL   = RGBColor(0xFF, 0xF1, 0x76)

# phase color ramp: prep is grey, then intensifying blues
RAMP = [GREYC, CYAN, MED, BLUE, NAVY]

L, R = 0.61, 12.73          # content margins
W = R - L                   # 12.12


def tbox(s, x, y, w, h, wrap=True):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return b


def put(p, runs, ls=1.0, align=None, space_before=None):
    p.line_spacing = ls
    if align:
        p.alignment = align
    if space_before is not None:
        p.space_before = Pt(space_before)
    for text, size, bold, color, italic in runs:
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size); f.bold = bold; f.italic = italic
        f.color.rgb = color; f.name = "Arial"
    return p


def rect(s, x, y, w, h, fill=None, line=None, lw=0.75, rounded=False, adj=None):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded and adj is not None:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def chevron(s, x, y, w, h, fill, first=False):
    """Right-pointing process arrow. PENTAGON for the head, CHEVRON for the rest
       so they interlock like Cetera's staircase."""
    shape = MSO_SHAPE.PENTAGON if first else MSO_SHAPE.CHEVRON
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        sp.adjustments[0] = min(0.5, (h * 0.42) / w)
    except Exception:
        pass
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def badge(s, cx, cy, num, fill, d=0.26, size=9.5, tcol=WHITE):
    """Numbered circle used to tie a chevron to its activity column."""
    b = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2),
                           Inches(d), Inches(d))
    b.fill.solid(); b.fill.fore_color.rgb = fill
    b.line.color.rgb = WHITE; b.line.width = Pt(1.0)
    b.shadow.inherit = False
    tf = b.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    put(tf.paragraphs[0], [(num, size, True, tcol, False)], align=PP_ALIGN.CENTER)
    return b


def bullet(s, x, y, w, lead, detail, size=8.5, lead_col=NAVY, col=BLACK,
           line_in=None, ls=0.92):
    """Hanging-indent bullet with a bold lead phrase. Height is measured, not guessed."""
    if line_in is None:
        line_in = size * ls * 1.20 / 72.0          # Arial line box at this size
    indent = 82296 / 914400.0                       # hanging indent eats into width
    full = f"{lead} {detail}".strip()
    n = est_lines(full, w - indent, size, bold_prefix_len=len(lead))
    h = n * line_in + 0.05
    b = tbox(s, x, y, w, h)
    p = b.text_frame.paragraphs[0]
    p.line_spacing = ls
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', '82296'); pPr.set('indent', '-82296')
    runs = [("• ", size, False, col, False), (lead, size, True, lead_col, False)]
    if detail:
        runs.append((" " + detail if not detail.startswith((",", ":", ";")) else detail,
                     size, False, col, False))
    put(p, runs, ls=ls)
    return b, h


_FONT_PATH = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
_FONT_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
_font_cache = {}


def _font(size, bold=False):
    key = (round(size * 4), bold)
    if key not in _font_cache:
        from PIL import ImageFont
        _font_cache[key] = ImageFont.truetype(_FONT_BOLD if bold else _FONT_PATH,
                                              max(1, int(round(size * 4))))
    return _font_cache[key]


def text_width(text, size, bold=False):
    """Width of `text` in inches, Arial-metric (Liberation Sans is metric-compatible)."""
    return _font(size, bold).getlength(text) / 4.0 / 72.0


def est_lines(text, w_in, size, bold_prefix_len=0, bold_size=None):
    """True word-wrap simulation: how many lines `text` takes in a w_in-wide box.
       bold_prefix_len chars at the start are measured with the bold face."""
    words = text.split()
    if not words:
        return 1
    bs = bold_size or size
    lines, cur, cur_w, consumed = 1, "", 0.0, 0
    space = text_width(" ", size)
    for wd in words:
        is_bold = consumed < bold_prefix_len
        ww = text_width(wd, bs if is_bold else size, bold=is_bold)
        consumed += len(wd) + 1
        add = ww if not cur else space + ww
        if cur and cur_w + add > w_in:
            lines += 1
            cur, cur_w = wd, ww
        else:
            cur = (cur + " " + wd) if cur else wd
            cur_w += add
    return lines


def col_geometry(weights, gap=0.10, left=L, total=W):
    """Return [(x, w), ...] for columns whose widths follow `weights`."""
    n = len(weights)
    usable = total - gap * (n - 1)
    unit = usable / sum(weights)
    out, x = [], left
    for wt in weights:
        w = wt * unit
        out.append((x, w))
        x += w + gap
    return out


def sticky(s, text, x=10.45, y=0.04, w=2.85, h=1.0, size=7.5):
    sp = rect(s, x, y, w, h, fill=YELL)
    sp.name = "ReviewSticky"
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    put(tf.paragraphs[0], [(text, size, False, BLACK, False)], ls=0.95)
    return sp


def wipe(s, keep_names):
    n = 0
    for sh in [x for x in s.shapes if x.name not in keep_names]:
        sh._element.getparent().remove(sh._element)
        n += 1
    return n


def set_title(s, title, subtitle=None):
    for sh in s.shapes:
        if sh.name in ("2. Slide Title", "Title 2", "Title 1"):
            tf = sh.text_frame
            p = tf.paragraphs[0]
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            for extra in list(tf.paragraphs[1:]):
                extra._p.getparent().remove(extra._p)
            r = p.add_run(); r.text = title
            break
    if subtitle is not None:
        for sh in s.shapes:
            if sh.name in ("3. Subtitle", "Subtitle 140"):
                tf = sh.text_frame
                p = tf.paragraphs[0]
                for r in list(p.runs):
                    r._r.getparent().remove(r._r)
                r = p.add_run(); r.text = subtitle
                break
