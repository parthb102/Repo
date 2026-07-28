#!/usr/bin/env python3
"""Merge: their 07/28 file is the base (layout, SteerCos, deliverables, appendix);
   our audited copy goes into their text boxes. Front-of-sentence bolding only,
   de-worded, every Megha sticky item preserved. GS slide cut, appendix flagged."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = "theirs.pptx"
DST = "merged_vg.pptx"
YELL = RGBColor(0xFF, 0xF1, 0x76)
BLACK = RGBColor(0, 0, 0)
MED = RGBColor(0x06, 0x79, 0xC3)
GREY = RGBColor(0x63, 0x66, 0x6A)

prs = Presentation(SRC)
slides = list(prs.slides)


def shape(s, name):
    hits = [sh for sh in s.shapes if sh.name == name]
    assert len(hits) == 1, (name, len(hits))
    return hits[0]


def rewrite(sh, items, size=None, debullet=False, spc=6):
    """Replace a text box's paragraphs with (lead, rest) bullets, keeping the
       box's own first-paragraph list formatting. Bold only on the lead."""
    tf = sh.text_frame
    p0 = tf.paragraphs[0]
    base = size
    if base is None:
        base = p0.runs[0].font.size.pt if (p0.runs and p0.runs[0].font.size) else 11
    tmpl = copy.deepcopy(p0._p.get_or_add_pPr())
    if debullet:
        for tag in ("a:buChar", "a:buAutoNum", "a:buNone", "a:spcBef",
                    "a:buFont", "a:buFontTx", "a:buSzPct", "a:buSzPts",
                    "a:buClr", "a:buClrTx"):
            for el in tmpl.findall(qn(tag)):
                tmpl.remove(el)
        tmpl.set("marL", "0"); tmpl.set("indent", "0")
        sb = tmpl.makeelement(qn("a:spcBef"), {})
        pts = sb.makeelement(qn("a:spcPts"), {"val": str(int(spc * 100))})
        sb.append(pts)
        tmpl.insert(0, sb)
        tmpl.append(tmpl.makeelement(qn("a:buNone"), {}))
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    if debullet:
        old = p0._p.find(qn("a:pPr"))
        if old is not None:
            p0._p.remove(old)
        p0._p.insert(0, copy.deepcopy(tmpl))
    for i, (lead, rest) in enumerate(items):
        if i == 0:
            para = tf.paragraphs[0]
        else:
            para = tf.add_paragraph()
            para._p.insert(0, copy.deepcopy(tmpl))
        r = para.add_run(); r.text = lead
        r.font.size = Pt(base); r.font.bold = True
        if rest:
            r2 = para.add_run(); r2.text = rest
            r2.font.size = Pt(base); r2.font.bold = False
    return sh


def sticky(s, text, x=10.45, y=0.02, w=2.85, h=1.15):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.name = "ReviewSticky"
    sp.fill.solid(); sp.fill.fore_color.rgb = YELL
    sp.line.fill.background(); sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]; p.line_spacing = 0.95
    r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(7.5); f.color.rgb = BLACK; f.name = "Arial"
    return sp


# ==================== SLIDE 2: framing table ====================
s = slides[1]

rewrite(shape(s, "TextBox 20"), [
    ("Establish what every technology dollar delivers today,",
     " where to rebalance to fund the future, and how capabilities compare with best in class"),
])
rewrite(shape(s, "TextBox 25"), [
    ("Identify where AI creates the greatest business value",
     " and redesign priority workflows around human and AI collaboration"),
])
rewrite(shape(s, "TextBox 4"), [
    ("Baseline and benchmark", " technology spend across applications, infrastructure, cybersecurity, and labor, comparing cost, productivity, and delivery maturity against peers"),
    ("Pressure test high-potential areas", " through deep dives (infrastructure, vendor contracts, developer productivity) and 5-10 leadership interviews"),
    ("Prioritize initiatives into a high-level roadmap", " with savings, owners, sequencing, and implementation implications"),
], size=11)
rewrite(shape(s, "TextBox 5"), [
    ("Baseline all business domains", " on workforce, cost, and process data, and size the initial AI value at stake in each"),
    ("Validate value with leaders", " and select the two priority domains on value, readiness, and feasibility"),
    ("Reimagine both domains end to end,", " defining future-state workflows, roles of people and AI agents, talent model, controls, and technology requirements"),
], size=11)
rewrite(shape(s, "TextBox 38"), [
    ("", "Celonis Process Intelligence"),
    ("", "McKinsey AI value-at-stake sizing tool"),
])
rewrite(shape(s, "TextBox 47"), [
    ("Current-state baselines", " and quantified AI value at stake across priority domains"),
    ("Two priority domains", " selected for reimagination"),
    ("Future-state workflows", " and enabling capabilities (talent, operating model, controls)"),
], size=11)
# title to a single line so the positioning line fits beneath it
tsh = shape(s, "2. Slide Title")
for p in tsh.text_frame.paragraphs:
    for r in p.runs:
        r.font.size = Pt(20)
# grow the objective cells a touch for the new copy
shape(s, "TextBox 20").height = Inches(1.20)
shape(s, "TextBox 25").height = Inches(1.20)
# positioning line under the title
sub = s.shapes.add_textbox(Inches(0.62), Inches(0.76), Inches(11.0), Inches(0.24))
tf = sub.text_frame; tf.word_wrap = True
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
p = tf.paragraphs[0]
r = p.add_run()
r.text = "Two tracks over the same 8 weeks: growth, effectiveness, and efficiency levers, not a cost exercise alone"
r.font.size = Pt(11); r.font.color.rgb = MED; r.font.name = "Arial"

sticky(s, "MERGED: kept this page's table layout from the 07/28 pass; ported the sharpened copy. Objectives front-bolded; Approach kept thematic, tightened to align with the p3/p4 plans (timing lives on those pages); 'AA tool' named client-readably; 'business aligned business' typo fixed; outputs now commit to two domains (this file said 1-2 in one cell and two in another); en dashes removed; holistic-levers positioning added under the title.")

# ==================== SLIDE 3: Track 1 workplan ====================
s = slides[2]
t = shape(s, "2. Slide Title").text_frame
for p in t.paragraphs:
    for r in p.runs:
        if r.text.rstrip().endswith(","):
            r.text = r.text.rstrip().rstrip(",")

rewrite(shape(s, "TextBox 12"), [
    ("Launch data requests", " and prep data for analysis (e.g., OPEX spend, IT taxonomies, workforce data)"),
    ("Finalize stakeholders and logistics:", " meetings, workshops, and the steering committee calendar"),
    ("Finalize surveys and respondents,", " ready to launch in week 1"),
    ("Develop initial hypotheses", " on where value sits"),
], debullet=True)
rewrite(shape(s, "TextBox 16"), [
    ("Baseline technology spend", " across applications, infrastructure, cybersecurity, and internal and external labor"),
    ("Assess and benchmark:", " cost and productivity (Tech Performance Assessment), delivery maturity (Technology Quotient), and engineering maturity (DVI)"),
    ("Hold working sessions with technology leaders", " to refine benchmark results and build early hypotheses on levers, gaps, and value at stake"),
], debullet=True)
rewrite(shape(s, "TextBox 20"), [
    ("Run detailed analyses and deep dives", " across all high-potential areas, e.g., infrastructure spend, vendor spend via line-by-line contract reviews, and developer productivity"),
    ("Conduct 5-10 interviews", " with business and technology leaders to pressure test opportunity areas"),
    ("Size each opportunity", " against addressable spend, with the actions to capture it"),
], debullet=True)
rewrite(shape(s, "TextBox 27"), [
    ("Shape recommendations", " into go-forward initiatives"),
    ("Detail the plan:", " initiatives to launch, estimated savings, potential owners, and high-level implementation implications (e.g., tooling)"),
    ("Develop the high-level roadmap:", " sequencing and dependencies"),
], debullet=True)
for p in shape(s, "TextBox 18").text_frame.paragraphs:
    for r in p.runs:
        if "Execution roadmap" in r.text:
            r.text = r.text.replace("Execution roadmap", "High-level roadmap")

sticky(s, "MERGED: kept the 07/28 layout (chevrons, SteerCos, deliverables row); replaced the pasted sticky text with client-ready copy. All of Megha's items are here; typos fixed ('to, pressure test', 'area/e.g.'); leads front-bolded; sub-bullets flattened. 'Execution roadmap' in the wks 7-8 deliverable -> 'High-level roadmap' per Megha's own qualifier. Tech Debt Index deep dive NOT re-added, matching this file's asset list.")

# ==================== SLIDE 4: Track 2 workplan ====================
s = slides[3]
rewrite(shape(s, "TextBox 2"), [
    ("Inventory the business domains", " for impact sizing and prioritization"),
    ("Collect available domain data:", " workforce and headcount, cost, volume, decision points, and hand-offs"),
    ("Confirm stakeholders and logistics:", " meetings, workshops, and the steering committee"),
], debullet=True)
rewrite(shape(s, "TextBox 6"), [
    ("Confirm the domains in scope", " where material value is at stake, anchored in strategic priorities and value pools"),
    ("Build a detailed baseline per domain:", " process flow, personnel, and performance (e.g., FTE utilization, SLA attainment)"),
    ("Size the AI value at stake,", " using our sizing tool for an initial view by domain"),
], debullet=True)
rewrite(shape(s, "TextBox 10"), [
    ("Validate the value at stake", " in working sessions with business, technology, finance, and functions leaders"),
    ("Build the prioritization framework:", " value potential (cost), strategic readiness (data, talent), and feasibility (decision complexity, risk and regulatory scrutiny)"),
    ("Score domains on a heatmap scorecard", " and select the two priority domains for reimagination"),
], debullet=True)
rewrite(shape(s, "TextBox 14"), [
    ("Reimagine each process end to end", " with Celonis Process Intelligence: future-state journey, AI and non-AI solutions, roles of people and AI agents, and exception paths"),
    ("Refine and finalize the blueprints", " in workshops with business leaders and subject matter experts"),
    ("Define the supporting implications:", " talent, operating model, data, technology, and controls"),
    ("Consolidate across both domains,", " identifying reusable components and sequencing with Track 1 initiatives"),
], debullet=True)

for p in shape(s, "TextBox 13").text_frame.paragraphs:
    for r in p.runs:
        r.text = r.text.replace("Executive selection of 1\u20132 domains", "Executive selection of the two priority domains")
        r.text = r.text.replace("1\u20132 domains", "the two priority domains")
for p in shape(s, "TextBox 11").text_frame.paragraphs:
    for r in p.runs:
        r.text = r.text.replace("value-at-stakes", "value at stake")

sticky(s, "MERGED: same treatment. 'hand-ffs' and other typos fixed; en dash removed; 'human and agent roles' -> 'roles of people and AI agents' (nobody at Vanguard should read 'agent' as a service agent); AA tool named client-readably. The wks 5-8 box covers both reimagination phases as designed here, closing with a consolidate bullet linking to Track 1. Page commits to two domains; your sticky said 'identifies 1-2', say the word to restore the optionality.")

# ==================== appendix ops ====================
sticky(slides[4], "CUT from this appendix: the 'Summary view of activities' page recycled from another client (it still said 'across GS' with a mid-November management-committee timeline). Recoverable from the 07/28 file if a re-skinned version is wanted.", h=0.95)

for idx in (7, 8, 9):
    s = slides[idx]
    lb = s.shapes.add_textbox(Inches(11.55), Inches(0.14), Inches(1.55), Inches(0.20))
    p = lb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = "ILLUSTRATIVE"
    f = r.font; f.size = Pt(10); f.bold = True; f.italic = True
    f.color.rgb = GREY; f.name = "Arial"

sticky(slides[8], "FLAG: the ACCE row ('identify bottom quartile of contributors') is sensitive in a client proposal; consider cutting the row or softening the language. $XX placeholders intentional? Page marked Illustrative.", y=0.35, h=0.95)

# delete the recycled GS slide (original slide 7, idx 6) last
sldIdLst = prs.slides._sldIdLst
el = list(sldIdLst)[6]
rId = el.get(qn('r:id'))
sldIdLst.remove(el)
prs.part.drop_rel(rId)

prs.save(DST)
print("merged saved:", DST, "slides now:", len(prs.slides._sldIdLst))
