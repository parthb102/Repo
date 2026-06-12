#!/usr/bin/env python3
"""Populate the TAM 2.0 tranching workbook with KTLO platforms.

Takes the user-edited workbook as the base: keeps the Instructions tab
verbatim, rebuilds Summary (same layout: table at row 4, Link column) and
the capability sheets with dynamically sized sections, preserves the
user's entered rows (notes enriched), maps every KTLO / Decommission
platform from the TAM repository extract to exactly one sub-capability,
and adds a Parking tab for POS software that has no home tab.
"""

import os

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SRC = ("/root/.claude/uploads/361d2622-d1b2-5153-8fcc-f2ea488e3e40/"
       "187fca8a-TAM_2.0_NonTarget_State_Tranche_vShare.xlsx")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "TAM_2.0_NonTarget_State_Tranche_vShare.xlsx")

NAVY = "1F4E79"
HEADER_TINT = "DDEBF7"
LINK = "0563C1"
T1_FILL = "CDE4F5"
T2_FILL = "FFF2CC"
T3_FILL = "FBE2D5"
GRAY_FILL = "D9D9D9"
GRAYED_ROW_FILL = "F2F2F2"
GRAYED_ROW_FONT = "A6A6A6"
NOTE_GRAY = "595959"

CIRCLED = ["①", "②", "③", "④", "⑤", "⑥", "⑦", "⑧", "⑨", "⑩",
           "⑪", "⑫", "⑬", "⑭", "⑮", "⑯", "⑰", "⑱", "⑲", "⑳",
           "㉑", "㉒", "㉓", "㉔", "㉕", "㉖", "㉗", "㉘", "㉙", "㉚",
           "㉛", "㉜", "㉝", "㉞", "㉟", "㊱", "㊲", "㊳", "㊴", "㊵",
           "㊶", "㊷", "㊸", "㊹", "㊺", "㊻", "㊼", "㊽", "㊾", "㊿"]

thin_gray = Side(style="thin", color="BFBFBF")
thin_blue = Side(style="thin", color="4472C4")
TABLE_BORDER = Border(left=thin_gray, right=thin_gray,
                      top=thin_gray, bottom=thin_gray)
MATRIX_BORDER = Border(left=thin_blue, right=thin_blue,
                       top=thin_blue, bottom=thin_blue)

BUCKETS = [
    ("High", "Low", "K", "O", "Tranche 1", T1_FILL),
    ("High", "Medium", "L", "P", "Tranche 2", T2_FILL),
    ("High", "High", "M", "Q", "Tranche 3", T3_FILL),
    ("Medium", "Low", "K", "R", "Tranche 1", T1_FILL),
    ("Medium", "Medium", "L", "S", "Tranche 2", T2_FILL),
    ("Medium", "High", "M", "T", None, GRAY_FILL),
    ("Low", "Low", "K", "U", None, GRAY_FILL),
    ("Low", "Medium", "L", "V", None, GRAY_FILL),
    ("Low", "High", "M", "W", None, GRAY_FILL),
]

TABLE_HEADERS = ["#", "Platform", "State", "Imperative to exit",
                 "Exit complexity", "Supporting facts", "Tranche (auto)"]

COL_WIDTHS = {"A": 5, "B": 26, "C": 17, "D": 16, "E": 15, "F": 44, "G": 14,
              "H": 2, "I": 4, "J": 9, "K": 26, "L": 26, "M": 26, "N": 2}

TAB_COLORS = {
    "Summary": "1F4E79", "Boarding": "2E75B6", "Single-in": "00B0F0",
    "Sales & Servicing": "7030A0", "Payment Processing": "C00000",
    "VAS": "ED7D31", "Unified Data and AI": "00B050",
    "Global Infrastructure": "808080", "Single-out": "BF8F00",
    "Parking": "A6A6A6",
}

CAPABILITIES = [
    ("Boarding", ["Boarding Experience", "Ingress", "Boarding Orchestration",
                  "Integration & Messaging", "Credit & Underwriting",
                  "CPQ / Quote-pricing"]),
    ("Single-in", ["Global API", "Payment Apps & SDKs", "Developer Experience",
                   "Hosted Solutions", "Unified Proposition"]),
    ("Sales & Servicing", ["CRM", "Servicing Tools", "Sales Automation",
                           "Risks"]),
    ("Payment Processing", ["Gateway", "Switch", "Clearing Back-end",
                            "Authorization Host", "Transaction monitoring",
                            "APMs/LPMs"]),
    ("VAS", ["Routing", "Networking Payment Tokens", "DCC eDCC", "Fraud",
             "Credential Management", "Gifts", "Managed Optimization",
             "3DS / exemptions", "FX"]),
    ("Unified Data and AI", ["Overall"]),
    ("Global Infrastructure", ["Data Centers", "Cloud", "Observability",
                               "Security", "AI Tools", "Design Systems"]),
    ("Single-out", ["Portal", "Reporting", "Billing", "Disputes", "Payouts",
                    "PFaaS Products"]),
    ("Parking", ["POS software (parked — not tranched for now)"]),
]

NT = "Non-target state"
KE = ("repository flags initial disposition Keep & Enhance — confirm "
      "with SPOC")

# (platform, state, imperative, complexity, supporting facts)
ROWS = {}

ROWS[("Payment Processing", "Gateway")] = [
    ("Global Gateway", "Target state", "", "", ""),
    ("Edge", NT, "High", "High",
     "Non-strategic but high volume — roughly 700M transactions and $50B "
     "processed annually across 80K MIDs, over half of contribution — with "
     "high feature complexity, surcharging built under prior constraints "
     "and about $7M in yearly run cost; structured multi-year migration, "
     "modernization and exit plan with stop-sell in around 24 months"),
    ("WP Total / Emboss", NT, "High", "Medium",
     "Non-strategic with moderate volume at 10–25% of contribution; "
     "stop-sell within six months while migration plans to target are "
     "built"),
    ("Innovo", NT, "Medium", "Low",
     "Non-strategic and low volume — about 18M transactions and $2B "
     "processed across 2.3K MIDs, roughly 1% of volume and $20M of revenue "
     "— with a small feature gap and about $350K in yearly run cost; "
     "immediate stop-sell within three months alongside migration "
     "planning"),
    ("EWay", NT, "Medium", "Medium",
     "Non-strategic with moderate volume at 10–25% of contribution; "
     "stop-sell within six months while migration plans to target are "
     "built"),
    ("Sterling", NT, "Medium", "Low",
     "Non-strategic and low volume — about 46M transactions and $1.8B "
     "processed across 4.7K merchants — and on the agreed decommission "
     "list; immediate stop-sell within three months alongside migration "
     "planning. The repository records Sterling Gateway and SNAP as a "
     "single platform — consider merging with the SNAP row"),
    ("Tron", NT, "Medium", "Low",
     "Non-strategic and low volume legacy eCommerce gateway for the "
     "Germany business; immediate stop-sell within three months alongside "
     "migration planning"),
    ("SNAP", NT, "Medium", "Low",
     "Non-strategic and low volume and on the agreed decommission list; "
     "immediate stop-sell within three months alongside migration "
     "planning. Likely the same platform as Sterling Gateway in the "
     "repository — confirm and merge if so"),
    ("GP WebPay", NT, "Medium", "Low",
     "Non-strategic and low volume in contribution terms — the repository "
     "shows about 63M transactions and $3.1B processed across 11K MIDs "
     "with around $800K in yearly run cost; immediate stop-sell within "
     "three months alongside migration planning, noting the repository's "
     "initial Keep & Enhance disposition"),
    ("Ezidebit", NT, "Medium", "Medium",
     "Non-strategic with moderate volume at 10–25% of contribution; "
     "stop-sell within six months while building migration plans to "
     "target; related Ezidebit platforms for boarding, reporting and "
     "settlement are captured in their own areas"),
    ("International Business Card (IBC)", NT, "", "",
     "Non-strategic legacy gateway serving US merchants; about 5.6M "
     "transactions and $688M processed across 2.2K MIDs; annual run cost "
     "near $450K; on the decommission path"),
    ("Transaction Express (TXP)", NT, "", "",
     "Non-strategic legacy TSYS gateway for US merchants; about 60M "
     "transactions and $7.8B processed across 19K MIDs; annual run cost "
     "near $1.1M including the ACI licence; on the decommission path"),
    ("Transaction Central (TC)", NT, "", "",
     "Non-strategic legacy TSYS gateway for US merchants; about 36M "
     "transactions and $4.8B processed across 43K MIDs; annual run cost "
     "near $1.1M including third-party support; on the decommission path"),
    ("DialPay 2.0", NT, "", "",
     "Non-strategic IVR authorization application for legacy TSYS "
     f"merchants; usage concentrated in older portfolios; {KE}"),
    ("EVO ACH", NT, "", "",
     "Non-strategic ACH platform supporting PayFabric, already in "
     "maintenance mode with Heartland ACH as the go-forward platform; "
     "migration path agreed"),
    ("Heartland Bill Pay", NT, "", "",
     "Bill-payment platform integrating billing ISVs into Heartland ACH "
     "and Portico; about 35M transactions and $4.5B processed across 6.4K "
     f"MIDs; around $1.2M in yearly third-party cost; {KE}"),
    ("Collonade ACH (Heartland ACH)", NT, "", "",
     "Described in the repository as the go-forward ACH platform, "
     "strategic and growing; about 16M transactions and $24B processed "
     "across 8K MIDs with annual run cost near $2.5M — confirm whether "
     "this row should be Target state"),
    ("Payments Business Platform (PBP)", NT, "", "",
     "Non-strategic bill-payments platform from the NCR acquisition; "
     "about 33M transactions and $14B processed across 3.4K MIDs; vision "
     "is migration to Heartland Bill Pay"),
    ("EVO One-Click WS", NT, "", "",
     "Non-strategic payment-link generator for premium LATAM merchants "
     f"with a small footprint of about 1K MIDs; {KE}"),
    ("GPI: Etrans/Storebuilder", NT, "", "",
     "Non-strategic public web service routing traffic into Innovo, ACH "
     "and gift processing; already in maintenance mode alongside the Edge "
     "estate"),
    ("MultiPay", NT, "", "",
     "Non-strategic EVO LATAM gateway spanning multiple countries and "
     "currencies; about 49M transactions and $8.2B processed across 2K "
     f"MIDs; {KE}"),
    ("IPG", NT, "", "",
     "Non-strategic EVO eCommerce gateway for the UK, Irish, Spanish, "
     "German and CEE markets; about 7.5M transactions and $320M processed "
     "across 3.5K MIDs; annual run cost near $2.6M; on the decommission "
     "path"),
    ("EDPS", NT, "", "",
     "Non-strategic Greek gateway serving eCommerce and POS for NBG Pay "
     "and other acquirers; roughly 15M transactions a month and $400M "
     "processed across 6.7K MIDs; repository flags Further Input — "
     "confirm with SPOC"),
    ("NestPay", NT, "", "",
     "Non-strategic third-party card-present gateway for the Polish "
     "market; about 2.5M transactions and $147M processed; annual cost "
     "near $140K; migration path agreed"),
    ("ClearOne", NT, "", "",
     "Non-strategic UniversalPay gateway serving integrated solutions in "
     "Iberia across 1.4K MIDs; to be decommissioned in favour of "
     "Comercia 2.0"),
    ("Redsys Cyberpack", NT, "", "",
     "Non-strategic resold Redsys eCommerce gateway offered by Comercia; "
     "to be decommissioned in favour of Addon 2.0"),
    ("Redsys Paygold", NT, "", "",
     "Non-strategic resold pay-by-link service built on the Cyberpac "
     "product; follows the Cyberpack decommission path"),
    ("Way2Pay / Collect", NT, "", "",
     "Non-strategic pay-by-link service integrated only with IPG; small "
     "team and footprint; integration to GP eCom already underway"),
    ("CyberSource (AP resold)", NT, "", "",
     "Resold Visa-owned gateway across Asia Pacific; about 88M "
     "transactions and $5B processed across 10.5K MIDs on per-transaction "
     f"fees; {KE}"),
    ("CyberSource (Erste JV)", NT, "", "",
     "White-label Visa-owned gateway for the Erste joint venture, still "
     f"in development phase; {KE}"),
    ("CyberSource (Greece transit)", NT, "", "",
     "Resold Visa-owned gateway serving the Greek transit project; very "
     "small footprint of about 300K transactions across 29 MIDs"),
    ("MiGS/MPGS (AP resold)", NT, "", "",
     "Resold Mastercard-owned gateway across Asia Pacific; about 11M "
     f"transactions and $980M processed across 1.4K MIDs; {KE}"),
    ("MPGS (NBG Pay)", NT, "", "",
     "White-label Mastercard gateway and main eCommerce solution for NBG "
     "Pay; about 11M transactions and $1.4B processed across 10K MIDs "
     f"with roughly $1.8M in annual fees; {KE}"),
    ("NMI (HPY)", NT, "", "",
     "Resold third-party gateway serving NA SMB sales initiatives; "
     "sustain in KTLO per the repository"),
    ("PayPal (PayFlow Pro & Link)", NT, "", "",
     "Resold third-party gateway serving NA SMB sales initiatives; "
     "sustain in KTLO per the repository"),
]

ROWS[("Payment Processing", "Authorization Host")] = [
    ("GSAP", NT, "", "",
     "Strategic international authorization and draft-capture platform "
     "per the repository, with NA, EU and AP instances processing about "
     "2.7B transactions and $190B annually at roughly $8M in engineering "
     "and support cost; flagged Keep & Enhance, so confirm whether this "
     "row should be Target state"),
    ("Exchange & Exchange Lite", NT, "", "",
     "Described in the repository as a US go-forward frontend under the "
     "prior TAM; about 3.1B transactions and $108B processed across 150K "
     "MIDs; flagged Keep & Enhance, so confirm whether this row should "
     "be Target state"),
    ("GNAP", NT, "", "",
     "Non-strategic Canadian authorization host and the only host "
     "certified for Interac; about 2.6B transactions and $113B processed "
     "across 97K MIDs with annual cost near $7M; sunset depends on GSAP "
     "feature parity and the new Interac service"),
    ("NGTrans", NT, "", "",
     "Non-strategic EVO frontend authorization platform in the US; about "
     "232M transactions and $25B processed across 39K MIDs; repository "
     "flags Keep & Enhance with selective migration — confirm with SPOC"),
    ("NPP/MESA", NT, "", "",
     "Non-strategic processing host from the NCR acquisition being "
     "integrated to Sierra; about 210M transactions and $22B processed in "
     "the first half of 2024; around $1.6M in yearly engineering cost"),
    ("NWS", NT, "", "",
     "Non-strategic petro frontend for network services; about 1.1B "
     "transactions and $68B processed across 22K MIDs; annual run cost "
     "near $2M on aging COBOL infrastructure; maintain-only under the "
     "prior TAM"),
    ("VAPS", NT, "", "",
     "Non-strategic petro frontend for value-added processing; about "
     "3.5B transactions and $129B processed across 26K MIDs; annual run "
     "cost near $1.9M on aging COBOL infrastructure; maintain-only under "
     "the prior TAM"),
    ("GPE ACI Base24", NT, "", "",
     "Non-strategic GPE authorization host in Prague serving CEE markets "
     "and significant indirect business; about 2.6B transactions and "
     f"$82B processed; annual cost near $4.5M; {KE}"),
    ("eService ACI Postillion Real-Time", NT, "", "",
     "Non-strategic eService authorization platform covering Poland, UK, "
     "Ireland, Spain and Germany; about 3.6B transactions processed in "
     f"2022; annual run cost near $2.9M; {KE}"),
    ("WLP", NT, "", "",
     "Non-strategic legacy component processing PayOne traffic for the "
     "Germany business; about 24M transactions and $1.6B processed across "
     "9.6K MIDs; on the decommission path"),
    ("Way4", NT, "", "",
     "Non-strategic authorization and processing platform for NBG Pay in "
     "Greece; to be replaced by e-service systems after migration"),
    ("Omnipay", NT, "", "",
     "Non-strategic third-party Fiserv platform with separate PayPal and "
     "GP instances covering Europe and Australia; migration project "
     "currently on hold"),
    ("BAS", NT, "", "",
     "Non-strategic batch authorization platform for card-not-present and "
     "recurring portfolios; about 44M transactions and $10.5B processed "
     "across 11K merchants; annual cost near $500K"),
]

ROWS[("Payment Processing", "Clearing Back-end")] = [
    ("GMAS", NT, "", "",
     "Described in the repository as the international go-forward backend "
     "under the prior TAM; about 4.9B transactions and $800B processed "
     "across 1.3M merchants with $7.5M in yearly Endava cost; flagged "
     "Keep & Enhance with selective migration, so confirm whether this "
     "row should be Target state"),
    ("A360", NT, "", "",
     "Non-strategic EVO North America settlement and clearing backend; "
     "about 268M transactions and $28B processed across 54K MIDs; annual "
     "run cost near $1M; migrates to Passport or TMAS under Falcon"),
    ("MCC", NT, "", "",
     "Non-strategic Comercia legacy backend supporting clearing, "
     "settlement and wider backend processes for roughly 700K active "
     "MIDs; maintained for CaixaBank customers with a reduced process "
     "footprint while decommissioning proceeds"),
    ("DOMEX", NT, "", "",
     "Non-strategic Comercia application funding non-CaixaBank customers, "
     "interlinked with MCC processes; follows the MCC decommission path"),
    ("AS400", NT, "", "",
     "Non-strategic core acquiring system for NBG Pay spanning boarding, "
     "data management and clearing; to be replaced by e-service systems "
     "after migration"),
    ("Postillion Back-Office", NT, "", "",
     "Non-strategic eService clearing and data-capture back office in "
     "Poland feeding ZSFK, risk and reconciliation from a single data "
     f"centre with no disaster recovery; annual cost near $1.8M; {KE}"),
    ("ZSFK", NT, "", "",
     "Non-strategic merchant settlement system for European EVO countries "
     "outside Germany; about 3.4B transactions and $78B processed across "
     f"206K MIDs; knowledge transfer to eService underway; {KE}"),
    ("Highlander", NT, "", "",
     "Non-strategic Oracle EBS settlement platform for Australia and New "
     f"Zealand across roughly 30K MIDs; about $1.6M in yearly cost; {KE}"),
    ("eBalance EU", NT, "", "",
     "Non-strategic EVO settlement balance reconciliation system across "
     "European markets; about $400K in yearly cost; on the decommission "
     "path"),
    ("Oracle Settlement (R12)", NT, "", "",
     "Non-strategic legacy Oracle settlement application already "
     "transitioning to Oracle Cloud across US, Canada, UK and AP"),
    ("Amatl Digital Voucher Vault", NT, "", "",
     "Non-strategic merchant digital voucher service for EVO LATAM across "
     "10K MIDs; hosting cost near $270K; confirmed for decommission"),
]

ROWS[("Payment Processing", "APMs/LPMs")] = [
    ("DelayPay", NT, "", "",
     "Third-party BNPL and consumer financing service on the Edge gateway "
     "with negligible current usage; repository flags Further Input — "
     "confirm with SPOC"),
    ("MTS", NT, "", "",
     "Non-strategic money transfer application down to a single customer "
     "with the contract already ended; confirmed for decommission"),
]

ROWS[("Single-out", "Disputes")] = [
    ("DiMS", NT, "", "",
     "Customer-facing and internal disputes platform for NA and "
     "international merchants handling about 1.5M disputes a year at "
     "roughly $3.3M in annual cost; repository flags initial disposition "
     "Keep & Enhance — under TAM 2.0 the long-term north star is WDP, so "
     "confirm timing with SPOC"),
    ("E360", NT, "", "",
     "Non-strategic legacy EVO disputes platform across US and EU "
     "portfolios; about 82K disputes and $23M handled across 13K MIDs; "
     "around $610K in yearly cost; migration path agreed"),
    ("Interseptas", NT, "", "",
     "Non-strategic disputes platform for NBG Pay in Greece; replaced by "
     "E360 after migration with legacy-transaction disputes remaining "
     "until run-off"),
    ("TCO & GIT", NT, "", "",
     "Non-strategic Comercia legacy applications supporting dispute "
     "documentation and chargeback management in Iberia; on the "
     "decommission path"),
]

ROWS[("VAS", "Fraud")] = [
    ("eSafe", NT, "", "",
     "Non-strategic post-authorization risk monitoring for EVO merchants "
     "across LATAM, EU and US with about $150B monitored across 550K "
     "MIDs; US portfolio already migrating to RiskNet V8; annual cost "
     "near $1.1M"),
    ("RiskNet", NT, "", "",
     "Third-party risk operations platform for NA and AP direct "
     "portfolios monitoring about $760B across 1.5M MIDs; annual cost "
     "near $4.4M including the ACI licence; repository flags Further "
     "Input — confirm with SPOC"),
    ("GPE PRM", NT, "", "",
     "Non-strategic ACI fraud detection and prevention system for CEMA "
     "acquiring and issuing, also used by external institutions and for "
     f"TRA exemptions; {KE}"),
    ("Riskshield", NT, "", "",
     "Non-strategic fraud management platform for NBG Pay in Greece; to "
     "be replaced by eSafe after migration"),
    ("RiskLab", NT, "", "",
     "Non-strategic Heartland automated risk edits suspending "
     "transactions for review across roughly 165K active MIDs; costs "
     f"bundled with Passport; {KE}"),
    ("CyberSource Decision Manager", NT, "", "",
     "Resold Visa fraud decisioning used in UK and EU; sustain in KTLO "
     "per the repository"),
    ("Heartland Secure", NT, "", "",
     "Encryption and breach-protection program for SMB merchants; "
     "overlaps with P2PE ownership and needs strategy validation per the "
     "repository"),
    ("ProPay ARM (with Tableau)", NT, "", "",
     "Automated risk management deeply embedded in the ProPay payfac "
     f"platform with Tableau reporting, non-extensible beyond ProPay; {KE}"),
]

ROWS[("Boarding", "Boarding Experience")] = [
    ("Launchpad", "Target state", "", "", ""),
    ("ELAPP", NT, "", "",
     "Non-strategic TSYS Broomfield boarding application covering about "
     "178K merchants; around $835K in yearly cost; replaced by Base CRM "
     "and IDGen under ECHO"),
    ("Merchant Application Package (MAP)", NT, "", "",
     "Non-strategic Omaha boarding solution also used by CPay covering "
     "about 87K merchants; around $430K in yearly third-party cost"),
    ("Central Station (CPay)", NT, "", "",
     "Non-strategic all-in-one business management platform for the "
     "Central Payments portfolio across 30K MIDs spanning sales, "
     "underwriting, boarding and service; about $360K in yearly cost"),
    ("EVONow", NT, "", "",
     "Non-strategic partner and merchant online application feeding the "
     "underwriting system; about $150K in yearly cost; folds into the "
     "EVO simplification program"),
    ("MOSS CA", NT, "", "",
     "Non-strategic boarding platform for the legacy GP Canada market "
     f"with 2.3K active users; about $410K in yearly cost; {KE}"),
    ("MOSS AP", NT, "", "",
     "Non-strategic boarding platform for the legacy GP Asia Pacific "
     f"market with about 230 active users; about $460K in yearly cost; "
     f"{KE}"),
    ("MOSS UK", NT, "", "",
     "Non-strategic boarding platform for the legacy GP UK market with "
     f"about 1.1K active users; about $990K in yearly cost; {KE}"),
    ("MOSS Mexico", NT, "", "",
     "Non-strategic boarding platform for the LATAM market; limited "
     "usage detail in the repository"),
    ("i-Apply", NT, "", "",
     "Non-strategic application platform for NBG Pay serving all "
     "application types; to be replaced by e-service systems"),
    ("MVSI Onboard", NT, "", "",
     "Non-strategic onboarding platform for a subset of Ezidebit "
     "clients; about $50K in yearly licence cost; confirmed for "
     "decommission"),
    ("Online Boarding Portal", NT, "", "",
     "Non-strategic legacy Heartland and Portico boarding front end for "
     "partners feeding Instant Boarding"),
    ("Rapid Activation", NT, "", "",
     "Non-strategic legacy GPI boarding for partners and sales users "
     "with modest monthly boarding volume"),
    ("Hermes", NT, "", "",
     "Non-strategic third-party boarding system for Czech bank advisors "
     "and sales reps; about 26K clients processed; modest annual cost"),
    ("Portico Website (POS Gateway)", NT, "", "",
     "Non-strategic internal UI for manual Portico boarding, "
     "configuration and research across US, Canada and UK portfolios"),
    ("NetSuite (Cayan)", NT, "", "",
     "Non-strategic SaaS boarding and servicing customization for the "
     "Cayan portfolio with about 670 users; around $670K in yearly cost"),
    ("Payroll CRM", NT, "", "",
     "Non-strategic boarding system for payroll clients with about 150 "
     "users; modest annual cost"),
    ("eContracts", NT, "", "",
     "Non-strategic merchant agreement and eSignature generation used in "
     "Poland; about 78K merchants processed; modest annual cost"),
    ("AssureSign", NT, "", "",
     "Third-party eSignature vendor for ELAPP priced per document; "
     "follows the ELAPP replacement path"),
    ("Agreement Express", NT, "", "",
     "Third-party eSignature tool used by the Cayan NetSuite boarding "
     "flow; on the decommission path"),
    ("Switch", NT, "", "",
     "Non-strategic boarding system for the Mainstream portfolio "
     "acquisition processing on the legacy TSYS stack; confirmed for "
     "decommission"),
    ("Fi911", NT, "", "",
     "Third-party boarding and deployment platform supporting the NCR "
     "portfolio and Jumpman; about $240K in yearly cost; strategy is to "
     "move to in-house boarding"),
]

ROWS[("Boarding", "Ingress")] = [
    ("Instant Boarding", NT, "", "",
     "Service-based boarding interface processing about 32K US payments "
     "and payroll applications a year from external systems; repository "
     "flags initial disposition Keep & Enhance — confirm with SPOC"),
    ("ELAPP Boarding API", NT, "", "",
     "Non-strategic merchant boarding API for partner integrators with "
     "about 9.3K MIDs boarded; around $280K in yearly cost; replaced by "
     "Base CRM"),
    ("MerchantFlo", NT, "", "",
     "Non-strategic boarding integration with a B2B API into OnTrak and "
     "a legacy Mindbody email flow; about $130K in yearly cost; B2B is "
     "replaced by Base CRM while the Mindbody flow persists"),
    ("Automated Boarding", NT, "", "",
     "Automated merchant and terminal boarding for OpenEdge and Genius "
     f"on Edge; {KE}"),
    ("Multi Location Boarding", NT, "", "",
     "Non-strategic web application boarding multiple locations at once "
     "into Instant Boarding with about 1.8K users; modest annual cost"),
    ("GPI: EdgeBuilder", NT, "", "",
     "Non-strategic service boarding merchants to the Edge gateway and "
     "linking Salesforce data; already in maintenance mode"),
]

ROWS[("Boarding", "Boarding Orchestration")] = [
    ("Onboard", NT, "", "",
     "Non-strategic boarding system for the EVO frontend boarding into "
     "A360 and GSAP with Discover MAP and Amex Opt Blue enrollments; "
     "about $150K in yearly cost; on the decommission path"),
    ("Merchant Enrollment Advantage (MEA)", NT, "", "",
     "Non-strategic GP US boarding application orchestrating boarding "
     "across host systems for about 113K merchants; around $260K in "
     "yearly cost"),
    ("MMS", NT, "", "",
     "Non-strategic legacy front-end loader pushing bank and merchant "
     "information into TSYS systems across 3.9M merchants; about $490K "
     "in yearly cost; profiles still required in TMAS for full-service "
     "clients"),
    ("Discover Enrollment System (DES)", NT, "", "",
     "Non-strategic GP application enrolling merchants with Discover "
     "across US and Canada portfolios"),
    ("Genesis", NT, "", "",
     "Non-strategic boarding system for NWS and VAPS merchants outside "
     "Client Manager with about 90 users; around $330K in yearly cost; "
     "repository flags Keep & Enhance with selective migration — confirm "
     "with SPOC"),
    ("Client Manager", NT, "", "",
     "Boarding system for Passport backend merchants with about 1.3K "
     "users; repository flags Keep & Enhance with selective migration — "
     "confirm whether this row should be Target state"),
    ("SEAT", NT, "", "",
     "Non-strategic sales entity and application workflow configuration "
     "tool with about 150 users; around $150K in yearly cost; repository "
     "flags Further Input — confirm with SPOC"),
    ("Flex (EVO boarding)", NT, "", "",
     "Non-strategic EVO boarding of demographics, plans and fees with "
     "equipment ordering and deployment ticketing; about $150K in yearly "
     "cost; on the decommission path. Distinct from the Canadian FLEX "
     "terminal application parked separately"),
]

ROWS[("Boarding", "Integration & Messaging")] = [
    ("The Central API", NT, "", "",
     "Communications backbone for Heartland POS providing boarding "
     "automation, invoicing, monitoring and service activations; about "
     "$1.3M in yearly cost; sustain in KTLO per the repository"),
]

ROWS[("Boarding", "Credit & Underwriting")] = [
    ("Underwriting Web", NT, "", "",
     "Non-strategic EVO tool for approving merchant applications with "
     "Mastercard MATCH and credit bureau checks; about $150K in yearly "
     "cost; repository flags Further Input — confirm with SPOC"),
    ("Universal Underwriting", NT, "", "",
     "Non-strategic service generating Experian reports and validation "
     "rules for Instant Boarding and Client Manager across about 32K "
     "applications a year; modest annual cost"),
    ("Mercury USA", NT, "", "",
     "Non-strategic credit, boarding and underwriting tool being sunset "
     "with EVONowHost taking over through the simplification program; "
     "noted as hard to scale with a history of security issues"),
    ("Jarvis", NT, "", "",
     "Non-strategic Get Beyond underwriting and merchant maintenance "
     "tool integrating with TSYS and Fiserv; about $150K in yearly cost; "
     "on the decommission path"),
]

ROWS[("Boarding", "CPQ / Quote-pricing")] = [
    ("Fee Rule Manager", NT, "", "",
     "Non-strategic pricing manipulation tool for Heartland merchants "
     "with about 120 users; repository flags Keep & Enhance with "
     "selective migration — confirm with SPOC"),
    ("Heartland Pricing Portal", NT, "", "",
     "Non-strategic pricing and retention tooling for Heartland "
     "merchants with about 220 users; repository flags Keep & Enhance "
     "with selective migration — confirm with SPOC"),
    ("Margin Service", NT, "", "",
     "Non-strategic margin calculator API for Heartland payments "
     "products; migration path agreed"),
]

ROWS[("Single-in", "Global API")] = [
    ("Comercia Addon2.0 API", NT, "", "",
     "Non-strategic Addon 2.0 integration API for Comercia merchants "
     "offered as JavaScript or host-to-host; on the decommission path"),
]

ROWS[("Single-in", "Payment Apps & SDKs")] = [
    ("Deployed Apps (XCharge, RCM)", NT, "", "",
     "Non-strategic legacy deployed applications tightly integrated with "
     "the Edge gateway; follow the Edge exit path"),
    ("Middleware for Aloha & Symphony POS", NT, "", "",
     "Non-strategic middleware between the gateway and Aloha or Symphony "
     "POS across 1.4K merchants; meaningful margin contribution of about "
     "$4.5M; sustain in KTLO per the repository"),
    ("Spidr", NT, "", "",
     "Non-strategic Windows middleware used mainly in parking across "
     "4.3K merchants; about 10M transactions and $150M processed; "
     "effectively no dedicated cost"),
    ("UK-Ingenico SDK", NT, "", "",
     "Non-strategic SDKs fronting legacy Ingenico applications for the "
     "UK ISV market; on the decommission path"),
    ("The Phoenix Group (TPG)", NT, "", "",
     "Third-party terminal deployment service sourcing hardware for "
     f"Omaha, EVO, Canada and CPay across about 13K devices; {KE}"),
    ("POS Portal", NT, "", "",
     "Third-party terminal deployment service for Cayan, Omaha and "
     f"Broomfield with about 100K devices shipped; {KE}"),
    ("EVO TMS", NT, "", "",
     "Non-strategic eService terminal management system supporting about "
     f"100K terminals across EU and UK; {KE}"),
    ("GPE TMS", NT, "", "",
     "Non-strategic GPE terminal management system covering about 91K "
     "terminals, tightly integrated with boarding; around $480K in "
     f"yearly cost; {KE}"),
    ("Terminal Management Services (TMS)", NT, "", "",
     f"Non-strategic terminal management supporting Genius across about "
     f"91K MIDs; {KE}"),
    ("Aurora TMS", NT, "", "",
     "Non-strategic LATAM platform administering remote terminal "
     "application loads for POS and iAcepta across 93K MIDs; about 166M "
     "transactions and $10.6B processed; around $180K in yearly cost"),
    ("WebTops", NT, "", "",
     "Non-strategic Heartland terminal configuration platform across "
     f"about 100K MIDs; around $1M in yearly cost; {KE}"),
    ("One Touch", NT, "", "",
     "Non-strategic internal UI managing terminal download files and "
     "parameters; no dedicated cost; on the decommission path"),
    ("AIMS", NT, "", "",
     "Non-strategic legacy Global US deployment system with about 190 "
     "users; around $130K in yearly cost; on the decommission path"),
    ("Deployment Manager", NT, "", "",
     "Non-strategic deployment web API used by ShipExec with warehouse "
     "printing services; modest annual cost; migration path agreed"),
    ("Inventory Manager", NT, "", "",
     "Non-strategic Canadian system tracking terminal physical locations "
     f"with about 160 users; around $130K in yearly cost; {KE}"),
    ("VeriCentre (Canada)", NT, "", "",
     "Non-strategic Canadian instance of the Verifone terminal "
     "management system tracking about 111K terminal IDs; sustain in "
     "KTLO per the repository"),
    ("GPC Injection System", NT, "", "",
     "Non-strategic key injection system using Verifone secure modules "
     "for Canadian hosts and gateways; sustain in KTLO per the "
     "repository"),
    ("Logistics", NT, "", "",
     "Non-strategic Heartland equipment tracking system with about 110 "
     "users; modest annual cost; migration path agreed"),
    ("Order Management System", NT, "", "",
     "Non-strategic merchant equipment ordering and configuration system "
     "with about 230 users; modest annual cost; migration path agreed"),
]

ROWS[("Single-in", "Developer Experience")] = [
    ("GPI Developer Portal", NT, "", "",
     "Non-strategic developer portal covering only GPI platforms with "
     "documentation and no sandbox; migration path agreed"),
    ("Heartland Developer Portal", NT, "", "",
     "Non-strategic developer portal for credentials, SDK libraries and "
     "documentation with about 18.5K users; repository flags Further "
     "Input — confirm with SPOC"),
    ("TSYS Developer Portal", NT, "", "",
     "Developer portal through which issuing clients and partners access "
     f"documentation and APIs; {KE}"),
    ("Comercia Developer Portal", NT, "", "",
     "Non-strategic documentation portal with SDKs and API explorer "
     "tools for the Spanish market"),
]

ROWS[("Sales & Servicing", "CRM")] = [
    ("Salesforce (TSYS Lightning)", NT, "", "",
     "Non-strategic instance managing ISV partner relationships for the "
     "Omaha portfolio with about 140 users; around $520K in yearly cost; "
     "migration path agreed"),
    ("Salesforce (Como)", NT, "", "",
     "Non-strategic instance providing lead management for Como clients "
     "with about 20 users; around $120K in yearly cost; repository flags "
     "Further Input — confirm with SPOC"),
    ("Salesforce (IBC)", NT, "", "",
     "Non-strategic instance for International Banc Card partners with "
     "about 120 users; around $120K in yearly cost; migration path "
     "agreed"),
    ("Salesforce (GetBeyond)", NT, "", "",
     "Non-strategic legacy GetBeyond instance with about 200 users; "
     "around $340K in yearly cost; migration path agreed"),
    ("Salesforce (Integrated)", NT, "", "",
     "Non-strategic instance covering the Integrated partner portfolio "
     "with about 1,000 users; around $2.8M in yearly cost; migration "
     "path agreed"),
    ("Salesforce (GetHired)", NT, "", "",
     "Non-strategic legacy training and servicing instance for the "
     "GetHired portfolio with about 15 users; around $90K in yearly "
     "cost; on the decommission path"),
    ("Salesforce (Springboard)", NT, "", "",
     "Non-strategic legacy Heartland Retail POS instance with about 30 "
     "users; around $110K in yearly cost; migration path agreed"),
    ("Salesforce (HPOS)", NT, "", "",
     "Instance managing the Heartland POS portfolio and dealer verticals "
     "with about 480 users; around $2M in yearly cost; repository flags "
     "Keep & Enhance with selective migration — confirm with SPOC"),
    ("Salesforce (International)", NT, "", "",
     "Lead management for international merchants across Canada, UK and "
     f"AP; around $3.7M in yearly cost; {KE}"),
    ("The Landing (Salesforce)", NT, "", "",
     "Non-strategic Get Beyond sales and service cloud with about 120 "
     "users; around $530K in yearly cost; on the decommission path"),
    ("Sugar CRM", NT, "", "",
     "Non-strategic sales pipeline tool managed separately per European "
     "country with about 135 users; around $110K in yearly cost; "
     "migration path agreed"),
    ("Korn Ferry Sell", NT, "", "",
     "Third-party account planning package inside Salesforce for large "
     "accounts with about 110 users; around $180K in yearly cost; on the "
     "decommission path"),
]

ROWS[("Sales & Servicing", "Servicing Tools")] = [
    ("OnTrak", NT, "", "",
     "Non-strategic Broomfield partner and merchant servicing suite "
     "including month-end residual processing of about $340M in rev "
     "share; around $2.1M in yearly cost; replaced by Base CRM and "
     "IDGen; repository flags Keep & Enhance with selective migration — "
     "confirm with SPOC"),
    ("Heartland Servicing Site", NT, "", "",
     "Non-strategic central hub for Heartland merchant servicing with "
     "about 2.5K users; modest annual cost; migration path agreed"),
    ("Document Management Service", NT, "", "",
     "Non-strategic document hub for Heartland merchant application "
     "uploads and downloads; modest annual cost; migration path agreed"),
    ("Finance Portal", NT, "", "",
     "Non-strategic internal app for modifying merchant banking "
     "information with about 190 users; modest annual cost; migration "
     "path agreed"),
    ("SGM", NT, "", "",
     "Non-strategic sales hierarchy tool assigning Heartland reps to "
     "sales areas with about 150 users; modest annual cost; migration "
     "path agreed"),
    ("Affiliate Atlas", NT, "", "",
     "Non-strategic sales affiliate portal for pricing updates with "
     "about 60 users; modest annual cost; migration path agreed"),
    ("Integrated Partner Service", NT, "", "",
     "Non-strategic external app identifying certified software "
     "solutions for processing; modest annual cost"),
    ("GPeCom GAdmin", NT, "", "",
     "Non-strategic UK servicing and administration tool for GP eCom; "
     "migration path agreed"),
    ("Tramites", NT, "", "",
     "Non-strategic call-centre application handling merchant claims and "
     "calls in LATAM; on the decommission path"),
    ("Debt$Net", NT, "", "",
     "Non-strategic collections software feeding FNBO and TSYS general "
     "ledgers; effectively no dedicated cost; migration path agreed"),
    ("eCommission", NT, "", "",
     "Non-strategic EVO US residual system across about 480K accounts; "
     "around $100K in yearly cost; on the decommission path"),
    ("Sales Compensation", NT, "", "",
     "Non-strategic Get Beyond seller compensation and portfolio health "
     "tool with about 50 users; around $120K in yearly cost; migration "
     "path agreed"),
]

ROWS[("Sales & Servicing", "Sales Automation")] = [
    ("Aim", NT, "", "",
     "Non-strategic sales scheduling and dialer tool serving the EVO "
     "direct team only; about $150K in yearly cost; confirmed for "
     "sunset"),
]

ROWS[("VAS", "Gifts")] = [
    ("Opticard / EVO Gift / Chockview", NT, "", "",
     "Non-strategic closed-loop gift platforms with gateway "
     "functionality to third-party processors; on the decommission "
     "path"),
    ("Global Gift App (Ingenico TETRA)", NT, "", "",
     "Non-strategic Ingenico-based Heartland gift payment application "
     "for Broomfield wholesale merchants; migration path agreed"),
    ("GPI: XGift Online Portal & Balance Check", NT, "", "",
     "Non-strategic merchant gift-card management and consumer "
     "balance-check web apps; already in maintenance mode"),
    ("HGM Tool", NT, "", "",
     "Non-strategic Windows module managing gift card products and SKUs "
     "with about 580 users; modest annual cost; migration path agreed"),
    ("Engage Loyalty Platform", NT, "", "",
     "Non-strategic loyalty platform powering ACE SMB, GP Marketplace "
     "and enterprise POS loyalty; around $1.5M in yearly cost across the "
     "Engage estate; migration path agreed"),
    ("Beanstalk (CRM + Loyalty)", NT, "", "",
     "Non-strategic customer engagement and loyalty cloud powering "
     "Whataburger and enterprise Xenial clients with about 15M end "
     "users; around $1.7M in yearly cost; migration path agreed"),
    ("GRS", NT, "", "",
     "Non-strategic legacy merchant loyalty platform on the NA frontend "
     "estate; kept lights-on under the prior TAM"),
    ("Micropayments Wave Rider", NT, "", "",
     "Non-strategic prepaid paydown front end for the gift platform in "
     "laundromat settings; sustain in KTLO per the repository"),
]

ROWS[("Unified Data and AI", "Overall")] = [
    ("TMS Data Warehouse", NT, "", "",
     "Non-strategic primary warehouse for Sierra and TMAS merchant data "
     "powering internal TSYS reporting; around $490K in yearly cost; "
     "migration path agreed"),
    ("Broomfield DataWarehouse", NT, "", "",
     "Non-strategic legacy reporting repository for Broomfield "
     "portfolios, gateways and OnTrak; around $450K in yearly cost; "
     "future direction is the Caspian datalake"),
    ("IBM Netezza Data Warehouse", NT, "", "",
     "Non-strategic data store for EVO LATAM supporting about 300K "
     "active MIDs; around $240K in yearly development cost; migration "
     "path agreed"),
    ("ETL implementation", NT, "", "",
     "Non-strategic home-grown transformation layer feeding internal and "
     "partner systems for Mexico and Chile; about 1.2B transactions and "
     "$64B flowing through; migration path agreed"),
    ("Big Data + HUE", NT, "", "",
     "Non-strategic NBG Pay data warehouse recalculating merchant "
     "service fees with HUE reporting; replaced after migration"),
    ("EDW", NT, "", "",
     "Non-strategic data warehouse and reporting system for NBG Pay; "
     "replaced after migration"),
    ("MNAC", NT, "", "",
     "Non-strategic data warehouse and reporting for internal Spanish "
     "management with about 10 direct users; confirmed for decommission"),
    ("Data Warehouse (Oceania)", NT, "", "",
     "Non-strategic local Oceania warehouse for internal and limited "
     "external use; migration path agreed"),
    ("Profisee (MDM)", NT, "", "",
     "Non-strategic MDM tool enabling real-time prospecting lookups for "
     "NA merchant sales; about $85K in yearly cost; migration path "
     "agreed"),
]

ROWS[("Single-out", "Portal")] = [
    ("Heartland InfoCentral", NT, "", "",
     "Non-strategic main reporting platform for US merchants with about "
     "19K monthly active users; functionality moves to My Account"),
    ("Open Edge View (OEV)", NT, "", "",
     "Non-strategic reporting portal for the Edge gateway with about "
     "35K unique monthly users; around $650K in yearly cost; "
     "functionality moves to My Account"),
    ("Merchant Success Hub", NT, "", "",
     "Non-strategic reporting portal for the Edge gateway with about "
     "14K unique monthly users; functionality moves to My Account"),
    ("Merchantware Portal", NT, "", "",
     "Non-strategic reporting portal for the Merchantware gateway; "
     "functionality moves to My Account alongside the other legacy "
     "portals"),
    ("Translink", NT, "", "",
     "Non-strategic reporting and servicing portal for legacy TSYS "
     "merchants with about 225K external users; around $370K in yearly "
     "cost; replaced by Base CRM under ECHO"),
    ("GPE Merchant Portal", NT, "", "",
     "Non-strategic main reporting platform for GPE merchants with "
     "about 89K external users; around $130K in yearly cost; replaced "
     "by Marketplace"),
    ("eService Client Portal", NT, "", "",
     "Non-strategic web portal for EVO merchants in Poland, Czechia and "
     "Hungary with about 180K users across multiple branded instances; "
     "migration path agreed"),
    ("BIS", NT, "", "",
     "Non-strategic main portal for EVO staff and merchants in Germany "
     "with about 5K users; under $50K in yearly cost; migration path "
     "agreed"),
    ("BRC", NT, "", "",
     "Non-strategic customer portal for EVO merchants in Spain, Ireland "
     "and UK with about 78K merchant accounts and limited ongoing "
     "development; migration path agreed"),
    ("MyEVO", NT, "", "",
     "Non-strategic merchant portal noted for low experience quality "
     "and high third-party build cost; about $150K in yearly cost; "
     "replaced by EVO in Control"),
    ("EVO Connect", NT, "", "",
     "Non-strategic EVO in Control portal due to be replaced by the "
     "partner portal; about $150K in yearly cost; repository flags "
     "Further Input — confirm with SPOC"),
    ("IPG Portal", NT, "", "",
     "Non-strategic gateway-level merchant portal for IPG; follows the "
     "IPG decommission path"),
    ("MyEway", NT, "", "",
     "Non-strategic main reporting platform for Eway clients in "
     "Oceania; sustain in KTLO per the repository"),
    ("EziOnline", NT, "", "",
     "Non-strategic main reporting platform for Ezidebit clients "
     "fronting Highlander; sustain in KTLO per the repository"),
    ("Partner Success Hub", NT, "", "",
     "Non-strategic Salesforce-communities partner portal for "
     "integrated partners; sustain in KTLO per the repository"),
    ("Heartland Partner Portal", NT, "", "",
     "Partner-facing portal for lead progression, referrals and "
     "residuals on a Snowflake source; repository flags Keep & Enhance "
     "with selective migration — confirm with SPOC"),
    ("PayPortal (Chargezoom)", NT, "", "",
     "Third-party B2B software with level 2 and 3 processing, invoicing "
     "and QuickBooks syncing for about 500 US SMBs; sustain in KTLO per "
     "the repository"),
    ("Beyond One", NT, "", "",
     "Non-strategic Get Beyond platform hosting the Sellers, Payments "
     "and Commerce apps for about 12.5K MIDs; around $285K in yearly "
     "cost; on the decommission path"),
    ("WinCOINS", NT, "", "",
     "Non-strategic legacy portal with limited repository detail; "
     "migration path agreed"),
]

ROWS[("Single-out", "Reporting")] = [
    ("Business Objects (SAP)", NT, "", "",
     "Non-strategic on-premise BI suite for Omaha partners covering "
     "about 242K MIDs; around $200K in yearly cost; migration path "
     "agreed"),
    ("Cognos (IBM)", NT, "", "",
     "Non-strategic on-premise BI suite for TSYS partners covering "
     "about 104K MIDs; around $690K in yearly licence cost; migration "
     "path agreed"),
    ("SSRS", NT, "", "",
     "Non-strategic server-based reporting for Heartland data sources "
     "with about 3.4K users; bundled with SQL Server licensing; "
     "migration path agreed"),
    ("Kepler", NT, "", "",
     "Non-strategic internal reporting for the Postillion authorization "
     "estate with about 500 users; modest annual cost; migration path "
     "agreed"),
    ("eService Oracle BI", NT, "", "",
     "Non-strategic Oracle BI management reporting for the eService "
     "business; migration path agreed"),
    ("Microsoft AX", NT, "", "",
     "Non-strategic internal reporting across LATAM, UK and EU EVO "
     "entities; migration path agreed"),
    ("Power BI (Oceania)", NT, "", "",
     "Internal Power BI reporting for Oceania; migration path agreed"),
]

ROWS[("Single-out", "Billing")] = [
    ("Netsuite (Eway)", NT, "", "",
     "Non-strategic billing and accounting for Eway eCommerce "
     "merchants; sustain in KTLO per the repository"),
    ("eStatement Portal", NT, "", "",
     "Non-strategic EVO US statement review, generation and branding "
     "system; around $50K in yearly cost; on the decommission path"),
    ("IRIS Billing", NT, "", "",
     "Non-strategic NA billing support application already moving to "
     "Oracle in July; around $60K in yearly licence cost"),
]

ROWS[("Single-out", "Payouts")] = [
    ("EMRS", NT, "", "",
     "Non-strategic EVO reserve, bank-reject and collections management "
     "reconciling merchant-level settlement; around $300K in yearly "
     "cost; on the decommission path"),
]

ROWS[("Parking", "POS software (parked — not tranched for now)")] = [
    ("Heartland Restaurant", NT, "", "",
     "SMB restaurant POS suite across about 12K MIDs and 50K stations; "
     f"around $730K in yearly AWS cost; {KE}"),
    ("Heartland Retail", NT, "", "",
     "Cloud-native SMB retail POS across about 1.1K MIDs; around $3.7M "
     f"in yearly cost; {KE}"),
    ("Cash Register Express (CRE)", NT, "", "",
     "Retail POS with prepaid paydown across about 2.2K MIDs and "
     "roughly $210M processed; sustain in KTLO per the repository"),
    ("Liquor Point of Sale (LPOS)", NT, "", "",
     "Liquor vertical POS across about 580 MIDs and roughly $103M "
     "processed; on the decommission path"),
    ("Dinerware", NT, "", "",
     "Windows restaurant POS hosted on merchant networks; about $310K "
     "in yearly AWS cost; on the decommission path"),
    ("Digital Dining", NT, "", "",
     "Windows restaurant POS hosted on merchant networks; on the "
     "decommission path"),
    ("Spectrum", NT, "", "",
     "US terminal payment application on Ingenico across about 65K MIDs "
     "and roughly $23B processed; on the decommission path"),
    ("HeartSIP", NT, "", "",
     "US terminal payment application on Ingenico across about 8K MIDs "
     "and roughly $13B processed; on the decommission path"),
    ("ConApp", NT, "", "",
     "Terminal payment application for petro and parking with about "
     "5.3K terminals; a new Android version exists"),
    ("FLEX (Canada terminal app)", NT, "", "",
     "Semi-integrated Ingenico Tetra payment application with about 8K "
     "active terminals and $0.5B processed; migration path agreed"),
    ("Apriva MPOS - US", NT, "", "",
     "Mobile POS suite for micro merchants; never had a platform home "
     "per the repository"),
    ("Mobile Inventory App (MIA)", NT, "", "",
     "Heartland mobile inventory companion for Cash Register Express; "
     "sustain in KTLO per the repository"),
    ("Heartland Mobile Manager (HMM)", NT, "", "",
     "POS-agnostic mobile monitoring app for Heartland Restaurant and "
     "CRE; around $430K in yearly cost; sustain in KTLO per the "
     "repository"),
    ("Diamond Cloud", NT, "", "",
     "Cloud and terminal client solution from EVO used in LATAM and EU; "
     "being demised under the QPR project"),
    ("Takepayments TP+", NT, "", "",
     "Terminal+ competitor in the UK boarding about 1.9K MIDs a month; "
     "needs fulfillment and tiered support solved before moving to "
     "Terminal+"),
    ("SmartPay", NT, "", "",
     "Resold integrated and standalone terminals for Ezidebit; about "
     "1.2M transactions and $127M processed; to be migrated to Quest"),
    ("Redsys (terminal app)", NT, "", "",
     "Payment application on Comercia legacy terminals; being demised "
     "in favour of MiTPV and Comercia 2.0"),
    ("Softpos/Likepos", NT, "", "",
     "Android softPOS payment app with about 5K users; around $33K in "
     "yearly cost; on the decommission path"),
    ("GPE Tap on Phone", NT, "", "",
     "Rubean-based softPOS solution for micro merchants and parcel "
     "delivery in CEE; on the decommission path"),
    ("EDPS Android Terminals", NT, "", "",
     "Android POS terminals for the Greek market; metrics tracked under "
     "the EDPS gateway line"),
    ("iAcepta", NT, "", "",
     "Mexican POS system across about 33K MIDs; around $490K in yearly "
     "cost; migration path agreed"),
]


def solid(color):
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def tranche_formula(r):
    return (
        f'=IF(OR($B{r}="",$D{r}="",$E{r}=""),"",'
        f'IF(OR($C{r}="Target state",$C{r}="Exception"),"Out of scope",'
        f'IF($D{r}="Low","No tranche",'
        f'IF($E{r}="Low","Tranche 1",'
        f'IF($E{r}="Medium","Tranche 2",'
        f'IF($D{r}="High","Tranche 3","No tranche"))))))'
    )


def helper_formula(r, imperative, complexity):
    return (
        f'=IF(AND($B{r}<>"",$C{r}<>"Target state",$C{r}<>"Exception",'
        f'$D{r}="{imperative}",$E{r}="{complexity}"),'
        f'$A{r}&"  "&$B{r},"")'
    )


def write_section(ws, sheet_name, s, subcap, dv_state, dv_score, seeds):
    n = max(10, len(seeds) + 4)
    r1, r2 = s + 2, s + 1 + n

    for col in "ABCDEFG":
        ws[f"{col}{s}"].fill = solid(NAVY)
    hc = ws[f"A{s}"]
    hc.value = subcap
    hc.font = Font(bold=True, size=11, color="FFFFFF")
    hc.alignment = Alignment(vertical="center", indent=1)
    ws.merge_cells(f"A{s}:F{s}")
    g = ws[f"G{s}"]
    g.value = f'=HYPERLINK("#\'{sheet_name}\'!A4","▲ tracker")'
    g.font = Font(size=9, color="FFFFFF", underline="single")
    g.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[s].height = 22

    for i, h in enumerate(TABLE_HEADERS):
        c = ws.cell(row=s + 1, column=1 + i, value=h)
        c.fill = solid(HEADER_TINT)
        c.font = Font(bold=True, size=9, color=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = TABLE_BORDER
    ws.row_dimensions[s + 1].height = 18

    for i in range(n):
        r = r1 + i
        ws.row_dimensions[r].height = 20
        a = ws[f"A{r}"]
        a.value = CIRCLED[i] if i < len(CIRCLED) else str(i + 1)
        a.alignment = Alignment(horizontal="center", vertical="center")
        a.font = Font(size=11, color=NAVY)
        ws[f"B{r}"].alignment = Alignment(vertical="center", indent=1,
                                          wrap_text=True)
        ws[f"B{r}"].font = Font(size=10)
        for col in "CDE":
            c = ws[f"{col}{r}"]
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.font = Font(size=10)
        f = ws[f"F{r}"]
        f.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        f.font = Font(size=9)
        gcell = ws[f"G{r}"]
        gcell.value = tranche_formula(r)
        gcell.font = Font(size=9, italic=True, color=NAVY)
        gcell.alignment = Alignment(horizontal="center", vertical="center")
        for col in "ABCDEFG":
            ws[f"{col}{r}"].border = TABLE_BORDER
        for imp, cpx, _mcol, hcol, _lbl, _fill in BUCKETS:
            ws[f"{hcol}{r}"] = helper_formula(r, imp, cpx)
        if i < len(seeds):
            name, state, imp, cpx, facts = seeds[i]
            ws[f"B{r}"] = name
            ws[f"C{r}"] = state
            if imp:
                ws[f"D{r}"] = imp
            if cpx:
                ws[f"E{r}"] = cpx
            if facts:
                ws[f"F{r}"] = facts
                ws.row_dimensions[r].height = max(
                    30, 12 * (len(facts) // 75 + 1))

    dv_state.add(f"C{r1}:C{r2}")
    dv_score.add(f"D{r1}:E{r2}")

    ws.conditional_formatting.add(
        f"A{r1}:G{r2}",
        FormulaRule(
            formula=[f'OR($C{r1}="Target state",$C{r1}="Exception")'],
            fill=PatternFill(start_color=GRAYED_ROW_FILL,
                             end_color=GRAYED_ROW_FILL, fill_type="solid"),
            font=Font(color=GRAYED_ROW_FONT),
        ),
    )

    t = ws[f"K{s}"]
    t.value = "Assessment Outcome"
    t.font = Font(bold=True, size=11, color=NAVY)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(f"K{s}:M{s}")
    for col, lbl in zip("KLM", ["Low", "Medium", "High"]):
        c = ws[f"{col}{s + 1}"]
        c.value = lbl
        c.font = Font(bold=True, size=9, color=NOTE_GRAY)
        c.alignment = Alignment(horizontal="center", vertical="center")

    b = n // 3
    rem = n % 3
    sizes = [b + (1 if rem > 0 else 0), b + (1 if rem > 1 else 0), b]
    top = s + 2
    for bi, band in enumerate(["High", "Medium", "Low"]):
        bottom = top + sizes[bi] - 1
        jc = ws[f"J{top}"]
        jc.value = band
        jc.font = Font(bold=True, size=10, color=NOTE_GRAY)
        jc.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(f"J{top}:J{bottom}")
        for imp, cpx, mcol, hcol, lbl, fill in BUCKETS:
            if imp != band:
                continue
            tj = f"_xlfn.TEXTJOIN(CHAR(10),TRUE,{hcol}${r1}:{hcol}${r2})"
            cell = ws[f"{mcol}{top}"]
            cell.value = f'="{lbl}"&CHAR(10)&{tj}' if lbl else f"={tj}"
            cell.font = Font(size=9, color="17375E" if lbl else "404040")
            cell.alignment = Alignment(wrap_text=True, vertical="top",
                                       horizontal="left", indent=1)
            for rr in range(top, bottom + 1):
                cc = ws[f"{mcol}{rr}"]
                cc.fill = solid(fill)
                cc.border = MATRIX_BORDER
            ws.merge_cells(f"{mcol}{top}:{mcol}{bottom}")
        top = bottom + 1

    ic = ws[f"I{s + 2}"]
    ic.value = "Benefits"
    ic.font = Font(bold=True, size=10, color=NOTE_GRAY)
    ic.alignment = Alignment(horizontal="center", vertical="center",
                             text_rotation=90)
    ws.merge_cells(f"I{s + 2}:I{s + 1 + n}")
    cx = ws[f"K{s + 2 + n}"]
    cx.value = "Complexity"
    cx.font = Font(bold=True, size=10, color=NOTE_GRAY)
    cx.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(f"K{s + 2 + n}:M{s + 2 + n}")

    return r1, r2, s + n + 6


def build_capability_sheet(wb, cap, subcaps):
    ws = wb.create_sheet(cap)
    ws.sheet_properties.tabColor = TAB_COLORS[cap]
    ws.sheet_view.showGridLines = False
    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w
    for col in "OPQRSTUVW":
        ws.column_dimensions[col].hidden = True

    title = ("Parking — POS software (not tranched for now)"
             if cap == "Parking" else f"{cap} — Platform Assessment")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=15, color=NAVY)
    ws["A2"] = '=HYPERLINK("#\'Summary\'!A1","← Back to Summary")'
    ws["A2"].font = Font(size=10, color=LINK, underline="single")
    note = ("Parked so nothing from the KTLO list is lost — re-home or "
            "score these later if needed."
            if cap == "Parking" else
            "Pick a State, score High / Medium / Low — the tranche, the "
            "matrix and the Summary update automatically.")
    ws["C2"] = note
    ws["C2"].font = Font(size=9, italic=True, color=NOTE_GRAY)

    th = ws["A4"]
    th.value = "Tracker — jump to a sub-capability:"
    th.font = Font(bold=True, size=10, color=NAVY)
    th.fill = solid(HEADER_TINT)
    th.alignment = Alignment(vertical="center", indent=1)
    for col in "BC":
        ws[f"{col}4"].fill = solid(HEADER_TINT)
    ws.merge_cells("A4:C4")

    dv_state = DataValidation(
        type="list",
        formula1='"Target state,Non-target state,Exception"',
        allow_blank=True)
    dv_score = DataValidation(
        type="list", formula1='"High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(dv_state)
    ws.add_data_validation(dv_score)

    sections = {}
    s = len(subcaps) + 6
    for i, subcap in enumerate(subcaps):
        ws[f"A{5 + i}"] = f"{i + 1}."
        ws[f"A{5 + i}"].alignment = Alignment(horizontal="right")
        ws[f"A{5 + i}"].font = Font(size=10, color=NOTE_GRAY)
        link = ws[f"B{5 + i}"]
        link.value = f'=HYPERLINK("#\'{cap}\'!A{s}","{subcap}")'
        link.font = Font(size=10, color=LINK, underline="single")
        seeds = ROWS.get((cap, subcap), [])
        r1, r2, nxt = write_section(ws, cap, s, subcap, dv_state, dv_score,
                                    seeds)
        sections[subcap] = (s, r1, r2)
        s = nxt
    return sections


def build_summary(wb, all_sections, spocs, titles):
    ws = wb.create_sheet("Summary")
    ws.sheet_properties.tabColor = TAB_COLORS["Summary"]
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 5, "B": 22, "C": 28, "D": 30, "E": 14, "F": 11,
                   "G": 11, "H": 11, "I": 10}.items():
        ws.column_dimensions[col].width = w

    ws["A1"] = titles[0]
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = titles[1]
    ws["A2"].font = Font(size=10, italic=True, color=NOTE_GRAY)

    hdr_row = 4
    headers = ["#", "Capability", "Sub-capability", "SPOCs",
               "Platforms listed", "Tranche 1", "Tranche 2", "Tranche 3",
               "Link"]
    hdr_fills = {"F": T1_FILL, "G": T2_FILL, "H": T3_FILL}
    for i, h in enumerate(headers):
        col = chr(ord("A") + i)
        c = ws[f"{col}{hdr_row}"]
        c.value = h
        if col in hdr_fills:
            c.fill = solid(hdr_fills[col])
            c.font = Font(bold=True, size=10, color=NAVY)
        else:
            c.fill = solid(NAVY)
            c.font = Font(bold=True, size=10, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)
        c.border = TABLE_BORDER
    ws.row_dimensions[hdr_row].height = 24

    r = hdr_row + 1
    idx = 1
    for cap, subcaps in CAPABILITIES:
        if cap == "Parking":
            continue
        group_start = r
        for subcap in subcaps:
            s, r1, r2 = all_sections[cap][subcap]
            ws[f"A{r}"] = idx
            ws[f"C{r}"] = subcap
            names = spocs.get(subcap, "")
            d = ws[f"D{r}"]
            d.value = names
            d.font = Font(size=9)
            d.alignment = Alignment(vertical="center", wrap_text=True,
                                    indent=1)
            ws[f"E{r}"] = f"=COUNTA('{cap}'!$B${r1}:$B${r2})"
            for col, lbl in zip("FGH",
                                ["Tranche 1", "Tranche 2", "Tranche 3"]):
                ws[f"{col}{r}"] = (
                    f"=COUNTIF('{cap}'!$G${r1}:$G${r2},\"{lbl}\")")
            ws[f"I{r}"] = f'=HYPERLINK("#\'{cap}\'!A{s}","Open →")'
            ws[f"I{r}"].font = Font(size=10, color=LINK, underline="single")
            ws[f"A{r}"].font = Font(size=9, color=NOTE_GRAY)
            ws[f"A{r}"].alignment = Alignment(horizontal="center",
                                              vertical="center")
            ws[f"C{r}"].font = Font(size=10)
            ws[f"C{r}"].alignment = Alignment(vertical="center")
            for col in "EFGH":
                ws[f"{col}{r}"].alignment = Alignment(horizontal="center",
                                                      vertical="center")
                ws[f"{col}{r}"].font = Font(size=10)
            ws[f"I{r}"].alignment = Alignment(horizontal="center",
                                              vertical="center")
            for col in "ABCDEFGHI":
                ws[f"{col}{r}"].border = TABLE_BORDER
            n_lines = names.count("\n") + 1 if names else 1
            ws.row_dimensions[r].height = max(18, 6 + 13 * n_lines)
            r += 1
            idx += 1
        b = ws[f"B{group_start}"]
        b.value = cap
        b.font = Font(bold=True, size=10, color=NAVY)
        b.alignment = Alignment(horizontal="left", vertical="center",
                                indent=1)
        for rr in range(group_start, r):
            ws[f"B{rr}"].fill = solid("EDF2F8")
        if r - 1 > group_start:
            ws.merge_cells(f"B{group_start}:B{r - 1}")

    last_data = r - 1
    ws[f"D{r}"] = "Total"
    ws[f"D{r}"].font = Font(bold=True, size=10, color=NAVY)
    ws[f"D{r}"].alignment = Alignment(horizontal="right", indent=1)
    for col in "EFGH":
        c = ws[f"{col}{r}"]
        c.value = f"=SUM({col}{hdr_row + 1}:{col}{last_data})"
        c.font = Font(bold=True, size=10, color=NAVY)
        c.alignment = Alignment(horizontal="center")
    for col in "ABCDEFGHI":
        ws[f"{col}{r}"].border = Border(top=Side(style="medium", color=NAVY))

    pk_sub = CAPABILITIES[-1][1][0]
    s, r1, r2 = all_sections["Parking"][pk_sub]
    pr = r + 1
    ws[f"C{pr}"] = '=HYPERLINK("#\'Parking\'!A1","Parking — POS software →")'
    ws[f"C{pr}"].font = Font(size=9, italic=True, color=LINK,
                             underline="single")
    ws[f"E{pr}"] = f"=COUNTA(Parking!$B${r1}:$B${r2})"
    ws[f"E{pr}"].font = Font(size=9, italic=True, color=NOTE_GRAY)
    ws[f"E{pr}"].alignment = Alignment(horizontal="center")
    ws[f"D{pr}"] = "not tranched"
    ws[f"D{pr}"].font = Font(size=9, italic=True, color=NOTE_GRAY)
    ws[f"D{pr}"].alignment = Alignment(horizontal="right", indent=1)

    ws.freeze_panes = f"A{hdr_row + 1}"


def main():
    wb = load_workbook(SRC)

    # carry over SPOC bullets and titles from the user's Summary
    src = wb["Summary"]
    spocs = {}
    for r in range(5, 48):
        sub = src.cell(row=r, column=3).value
        if sub:
            spocs[sub] = src.cell(row=r, column=4).value or ""
    titles = (src["A1"].value or "Platform Assessment — Summary",
              src["A2"].value or "Tranche roll-up by capability")

    for name in ["Summary", "Boarding", "Single-in", "Sales & Servicing",
                 "Payment Processing", "VAS", "Unified Data and AI",
                 "Global Infrastructure", "Single-out"]:
        del wb[name]

    all_sections = {}
    for cap, subcaps in CAPABILITIES:
        all_sections[cap] = build_capability_sheet(wb, cap, subcaps)
    build_summary(wb, all_sections, spocs, titles)

    order = ["Instructions", "Summary"] + [c for c, _ in CAPABILITIES]
    wb._sheets = [wb[n] for n in order]

    wb.save(OUT)
    total = sum(len(v) for v in ROWS.values())
    print(f"Wrote {OUT} — {total} platform rows seeded")


if __name__ == "__main__":
    main()
