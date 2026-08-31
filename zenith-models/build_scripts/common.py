#!/usr/bin/env python3
"""Shared assembly flow for a workstream workbook."""
import openpyxl
from framework import TEMPLATE_PATH
from builders import build_rate_card, build_cover
from builders2 import ModelBuilder, fill_wave, build_summary

STD_DEVIATIONS = [
    "WAVE 'Recurring Benefits Total' run-rate cell K88: standard template has =SUM(-K68,K76,K84), "
    "which negates revenue in the run-rate column only (annual columns sum straight). Fixed here to "
    "=SUM(K68,K76,K84) so revenue-bearing cases pass the template's own K-vs-2030 check.",
    "WAVE ROI cell G163: standard template divides the BASELINE column (I117) by one-time cost; "
    "fixed to use the run-rate NRB (K117). Both G163 (ROI) and G164 (IRR) are IFERROR-guarded so a "
    "blank or zero-cost state shows 'n/a' instead of DIV/0- and NUM-errors.",
    "WAVE label cells I64/I92 ('TTM TTM 6/1/26' in the source, caused by a text period in H18): "
    "now display the Section I baseline period directly (=$H$18).",
    "Template's illustrative prefills in rows 79/107 (I79=J21 link, J107=J26 link) neutralized on "
    "unused rows so phantom values cannot appear; the Section I generic 'xxxx' placeholder blocks "
    "(template rows 31-57) are cleared on initiative WAVE sheets.",
    "On populated WAVE rows the K 'Ann. Run Rate' cell links the model's run-rate cell directly "
    "instead of recomputing I-J / J-I (J is built from I and that same cell, so the arithmetic is "
    "identical); this keeps K bit-identical to the 2030 column so the R-column run-rate checks stay "
    "exact in Excel.",
    "Payback timeline (row 161 / G165): the template's chain formula stalls whenever a year's "
    "cumulative cash flow is exactly zero (its own example shows payback '0 yrs' despite a -6.2M "
    "cumulative in 2027). Replaced with a robust per-year form: 1 if cumulative CF < 0, the "
    "crossing-year fraction when it turns positive, else 0; G165 shows '>5' when payback is not "
    "reached inside the 2026-30 horizon.",
    "The reference 'Standard Template' tab at the end of this file carries ONLY the error-guard "
    "fixes (G163/G164 IFERROR wraps, robust payback row) so the file recalculates clean; every "
    "other cell is byte-identical to the uploaded template, including its original K88/I64 quirks "
    "and all check machinery (R1, CI1:CM1), WACC 9.5%, cash-basis CapEx treatment. The "
    "'Instructions' tab is as uploaded except its dead broken-reference (REF) hidden drop-down cells were cleared "
    "(broken in the source file).",
]

LINKING_NOTES = [
    "Each initiative has a '<code> Model' sheet (all logic and inputs) and a '<code> WAVE' sheet "
    "(the standard template). The WAVE sheet contains NO hardcoded numbers for this initiative — "
    "its benefit/cost rows link (green) to the Model sheet's feed rows in $M.",
    "Model sheets use the same time grid as WAVE: columns L:P = FY2026–2030 annual, columns T:CA = "
    "Jan-2026..Dec-2030 monthly. Annual cells are SUMIFS of the monthly grid, so the WAVE "
    "monthly-vs-annual checks (CI1:CM1) pass by construction.",
    "The 'Ann. Run Rate' (K) on feed rows is defined as the 2030 value, matching the WAVE rule "
    "'Annualized Run Rate (K) = 2030 Value (P)'. Keep ramps reaching 100% before 2030.",
    "Benefits ramp linearly from the go-live date over 'months to full run rate' — same convention "
    "as the Servicing S2 model. One-time costs spread evenly over each milestone's billing months.",
    "People costs: milestone labor = Team size x Billing months x blended monthly rate (Rate Card) x "
    "(1 + contingency). Capitalization: each milestone carries a Cap % (default from Rate Card "
    "policy) that splits labor into WAVE CapEx vs OTC rows. Same for non-labor one-time items.",
]


def assemble(out_path, meta, rc_defaults, specs, summary_title,
             template_path=TEMPLATE_PATH, cap_note=None):
    wb = openpyxl.load_workbook(template_path)
    if "Business Case_Example" in wb.sheetnames:
        del wb["Business Case_Example"]
    # sanitize inherited errors (documented on Cover):
    # 1. Instructions hidden drop-down block carries dead =#REF! formulas
    ins = wb["Instructions"]
    for r in range(176, 212):
        cc = ins[f"E{r}"]
        if isinstance(cc.value, str) and "#REF!" in cc.value:
            cc.value = None
    # 2. blank-state ROI/IRR errors on the pristine reference template
    st = wb["Standard Template"]
    st["G163"] = '=IFERROR(I117/-SUM(L149:P149),"n/a")'
    st["G164"] = '=IFERROR(IRR($L$159:$P$159, 0.2),"n/a")'
    st["L161"] = "=IF(L160<0,1,0)"
    for cur, prev in [("M", "L"), ("N", "M"), ("O", "N"), ("P", "O")]:
        st[f"{cur}161"] = (f"=IF({cur}160<0,1,IF(AND({prev}160<0,{cur}159>0),"
                           f"-{prev}160/{cur}159,0))")
    rc = build_rate_card(wb, "Rate Card", rc_defaults)
    pair_names = []
    for spec in specs:
        mb = ModelBuilder(wb, spec, rc)
        feed = mb.build()
        fill_wave(wb, spec, feed, mb.sym)
        pair_names += [spec["model_sheet"], spec["wave_sheet"]]
    build_summary(wb, specs, summary_title, cap_note=cap_note)
    build_cover(wb, meta)
    # order: Cover, Instructions, Rate Card, Summary, pairs..., Standard Template
    order = ["Cover", "Instructions", "Rate Card", "Summary"] + pair_names + ["Standard Template"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + \
                 [ws for ws in wb._sheets if ws.title not in order]
    wb["Standard Template"].sheet_properties.tabColor = "808080"
    wb["Instructions"].sheet_properties.tabColor = "808080"
    wb.active = 0
    wb.save(out_path)
    print("saved", out_path)
    return out_path
