#!/usr/bin/env python3
"""
Shared builder framework for Zenith workstream workbooks.

Each workstream file = copy of the WAVE template workbook (keeps 'Instructions'),
with per-initiative Model sheets + WAVE sheets (copies of 'Standard Template'),
a Cover, a Rate Card, and a Summary sheet.

Column conventions on Model sheets (mirrors WAVE):
  B..C labels | D value | E unit | F..I source/notes (F wide) | K run-rate |
  L:P annual 2026-2030 | T:CA monthly Jan-26..Dec-30
"""
import openpyxl
from copy import copy as _copy
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

# ---------------------------------------------------------------- constants
TEMPLATE_PATH = "abd10a23-Zenith_Business_Case_Template_by_Lever_1.xlsx"
FONT = "Calibri"

MON_FIRST = column_index_from_string("T")   # Jan-2026
MON_LAST = column_index_from_string("CA")   # Dec-2030
YR_FIRST = column_index_from_string("L")    # 2026
YEARS = [2026, 2027, 2028, 2029, 2030]

# colors (financial-model conventions)
BLUE = "FF0000FF"      # hardcoded input
BLACK = "FF000000"
GREEN = "FF008000"     # cross-sheet link
GREY_FILL = "FFD9D9D9"
DARKBLUE_FILL = "FF1F3864"
BAND_FILL = "FFDDEBF7"     # section band
INPUT_FILL = "FFFFFFCC"    # template's own yellow input fill
WHITE = "FFFFFFFF"

FMT_USD = '$#,##0;($#,##0);"-"'
FMT_USD2 = '$#,##0.00;($#,##0.00);"-"'
FMT_USDM = '#,##0.000;(#,##0.000);"-";@'   # template's $USD M format
FMT_PCT = '0.0%;(0.0%);"-"'
FMT_PCT0 = '0%;(0%);"-"'
FMT_NUM = '#,##0;(#,##0);"-"'
FMT_NUM1 = '#,##0.0;(#,##0.0);"-"'
FMT_DATE = 'mmm-yy'
FMT_X = '0.00"x"'

THIN = Side(style="thin", color="FFBFBFBF")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
B_BOT = Border(bottom=Side(style="medium", color="FF1F3864"))


def f(bold=False, color=BLACK, size=11, italic=False, name=FONT):
    return Font(name=name, bold=bold, color=color, size=size, italic=italic)


def fill(rgb):
    return PatternFill("solid", fgColor=rgb)


def col(idx):
    return get_column_letter(idx)


class SheetWriter:
    """Row-cursor based writer that tracks named row anchors."""

    def __init__(self, ws):
        self.ws = ws
        self.r = 1
        self.anchors = {}

    def mark(self, name, row=None):
        self.anchors[name] = row if row is not None else self.r
        return self.anchors[name]

    def a(self, name):
        return self.anchors[name]

    def skip(self, n=1):
        self.r += n

    def cell(self, r, c, v=None, font=None, num=None, fl=None, align=None,
             wrap=False, border=None, comment=None, indent=0):
        cc = self.ws.cell(row=r, column=c)
        if v is not None:
            cc.value = v
        cc.font = font or f()
        if num:
            cc.number_format = num
        if fl:
            cc.fill = fill(fl)
        al = Alignment(horizontal=align, vertical="center", wrap_text=wrap,
                       indent=indent)
        cc.alignment = al
        if border:
            cc.border = border
        if comment:
            cc.comment = Comment(comment, "Model")
        return cc

    # -- high-level helpers -------------------------------------------------
    def title_row(self, text, sub=None):
        self.cell(self.r, 2, text, f(bold=True, size=16, color=WHITE), fl=DARKBLUE_FILL)
        for c in range(3, 15):
            self.cell(self.r, c, None, fl=DARKBLUE_FILL)
        self.ws.row_dimensions[self.r].height = 24
        self.r += 1
        if sub:
            self.cell(self.r, 2, sub, f(italic=True, size=10, color="FF595959"))
            self.r += 1

    def section(self, label):
        r = self.r
        self.cell(r, 2, label, f(bold=True, size=12, color="FF1F3864"), fl=BAND_FILL)
        for c in range(3, 15):
            self.cell(r, c, None, fl=BAND_FILL)
        self.ws.row_dimensions[r].height = 18
        self.r += 1
        return r

    def input_cell(self, r, c, v, num=None, comment=None):
        return self.cell(r, c, v, f(color=BLUE), num=num, fl=INPUT_FILL,
                         border=BORDER_ALL, comment=comment)

    def formula_cell(self, r, c, v, num=None, link=False, comment=None):
        return self.cell(r, c, v, f(color=GREEN if link else BLACK), num=num,
                         comment=comment)

    def label(self, r, text, c=2, bold=False, italic=False, indent=0, wrap=False):
        return self.cell(r, c, text, f(bold=bold, italic=italic), wrap=wrap,
                         indent=indent)

    def note(self, r, text, c=6, width_cols=6):
        return self.cell(r, c, text, f(italic=True, size=9, color="FF595959"),
                         wrap=True)


def month_header_rows(sw, yr_row=None, hd_row=None):
    """Write year row (like WAVE row 1) and EOMONTH date header row on a model
    sheet. Returns (year_row, date_row). Also writes annual year labels L:P."""
    if yr_row is None:
        yr_row = sw.r
        sw.r += 1
    if hd_row is None:
        hd_row = sw.r
        sw.r += 1
    ws = sw.ws
    # annual labels
    sw.cell(hd_row, YR_FIRST - 1, "Run rate", f(bold=True, size=9), align="center")
    for i, y in enumerate(YEARS):
        sw.cell(hd_row, YR_FIRST + i, y, f(bold=True, size=9), align="center")
    # monthly: year helper + EOMONTH chain
    for i in range(60):
        c = MON_FIRST + i
        y = 2026 + i // 12
        sw.cell(yr_row, c, y, f(size=8, color="FF808080"), align="center")
    sw.cell(hd_row, MON_FIRST, "=DATE(2026,1,31)", f(bold=True, size=8),
            num=FMT_DATE, align="center")
    for i in range(1, 60):
        c = MON_FIRST + i
        sw.cell(hd_row, c, f"=EOMONTH({col(c-1)}{hd_row},1)",
                f(bold=True, size=8), num=FMT_DATE, align="center")
    return yr_row, hd_row


def annual_from_monthly(sw, r, yr_row, num=FMT_USD, k_formula=None):
    """L:P = SUMIFS over the monthly grid of row r keyed by year row."""
    for i, y in enumerate(YEARS):
        c = YR_FIRST + i
        sw.formula_cell(
            r, c,
            f"=SUMIFS($T{r}:$CA{r},$T${yr_row}:$CA${yr_row},{y})",
            num=num)
    if k_formula:
        sw.formula_cell(r, YR_FIRST - 1, k_formula, num=num)


def monthly_window_spread(sw, r, hd_row, total_ref, start_ref, months_ref,
                          num=FMT_USD):
    """T:CA = total/months inside [start, start+months-1] month window."""
    for i in range(60):
        c = MON_FIRST + i
        m = f"{col(c)}${hd_row}"
        sw.formula_cell(
            r, c,
            f"=IF(AND({m}>=EOMONTH({start_ref},0),"
            f"{m}<=EOMONTH(EDATE({start_ref},{months_ref}-1),0)),"
            f"{total_ref}/{months_ref},0)", num=num)


def copy_wave_sheet(wb, new_title):
    """Duplicate 'Standard Template' with full fidelity (CF, views, freeze)."""
    src = wb["Standard Template"]
    dst = wb.copy_worksheet(src)
    dst.title = new_title
    for cf in src.conditional_formatting:
        for rule in cf.rules:
            dst.conditional_formatting.add(str(cf.sqref), _copy(rule))
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.sheet_view.zoomScale = src.sheet_view.zoomScale
    dst.freeze_panes = src.freeze_panes
    dst.sheet_properties.tabColor = "1F3864"
    return dst


def wave_write(ws, ref, value, green=True, num=None):
    """Write a link/value into a WAVE sheet cell, keeping its style but
    marking links green / inputs blue."""
    cc = ws[ref]
    cc.value = value
    fnt = _copy(cc.font)
    is_formula = isinstance(value, str) and value.startswith("=")
    fnt.color = openpyxl.styles.colors.Color(
        rgb=(GREEN if (green and is_formula) else BLUE if not is_formula else BLACK))
    fnt.name = FONT
    cc.font = fnt
    if num:
        cc.number_format = num
    return cc
