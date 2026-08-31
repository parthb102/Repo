#!/usr/bin/env python3
"""High-level sheet builders: Cover, Rate Card, Model sheets, WAVE linking,
Summary. Driven by per-initiative spec dicts."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.datavalidation import DataValidation
from framework import (SheetWriter, month_header_rows, annual_from_monthly,
                       monthly_window_spread, copy_wave_sheet, wave_write,
                       f, fill, col, FONT, BLUE, BLACK, GREEN, GREY_FILL,
                       DARKBLUE_FILL, BAND_FILL, INPUT_FILL, WHITE,
                       FMT_USD, FMT_USD2, FMT_USDM, FMT_PCT, FMT_PCT0,
                       FMT_NUM, FMT_NUM1, FMT_DATE, FMT_X,
                       MON_FIRST, MON_LAST, YR_FIRST, YEARS, BORDER_ALL)

MILESTONES = ["Plan / Design", "Build / Test", "Pilot / Change Mgmt.", "Scale / Run"]
CONT_TABLE = [("Low", 0.10), ("Medium", 0.15), ("High", 0.20), ("Very high", 0.50)]


def _model_cols(ws):
    widths = {"A": 3, "B": 46, "C": 13, "D": 15, "E": 12, "F": 15, "G": 15,
              "H": 15, "I": 13, "J": 2, "K": 12, "L": 11, "M": 11, "N": 11,
              "O": 11, "P": 11, "Q": 2, "R": 60, "S": 2}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    for i in range(MON_FIRST, MON_LAST + 1):
        ws.column_dimensions[get_column_letter(i)].width = 9


# ---------------------------------------------------------------- Rate Card
def build_rate_card(wb, ws_name, defaults):
    """defaults: dict with keys:
       roles: [(role, count, annual_rate)], basis ('Yearly'|'Monthly'|'Hourly'),
       hourly_rate, hours_day, days_month, build_note, run_rate_annual,
       cap_pcts: {milestone: pct}, currency_note
    """
    ws = wb.create_sheet(ws_name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "70AD47"
    _model_cols(ws)
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 16
    ws.column_dimensions["I"].width = 60
    sw = SheetWriter(ws)
    sw.r = 2
    sw.title_row("RATE CARD — people-cost engine (shared by all initiative models)")
    sw.skip(1)
    sw.cell(sw.r, 2, "Blue cells with yellow fill = inputs you can change. Black = formulas. "
            "All initiative models pull the blended rate and policy tables from this sheet.",
            f(italic=True, size=10, color="FF595959"))
    sw.skip(2)

    # --- 1. rate basis ---------------------------------------------------
    sw.section("1. RATE BASIS")
    sw.skip(1)
    r = sw.r
    sw.label(r, "Rate basis for people costs", bold=True)
    sw.input_cell(r, 4, defaults.get("basis", "Yearly"),
                  comment="Yearly / Monthly / Hourly — drives which rate column is used")
    dv = DataValidation(type="list", formula1='"Yearly,Monthly,Hourly"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"D{r}")
    sw.mark("basis", r)
    sw.note(r, "Choose how role rates are quoted, then TYPE YOUR RATES in the matching column of the "
               "role table below (Yearly->col D, Monthly->col E, Hourly->col F). The other columns show "
               "derived equivalents until you overwrite them. Hourly => monthly = $/hr x hours/day x days/month.", c=6)
    sw.skip(2)
    r = sw.r
    sw.label(r, "Working hours per day")
    sw.input_cell(r, 4, defaults.get("hours_day", 8), num=FMT_NUM)
    sw.mark("hours_day", r)
    sw.label(r + 1, "Billable days per month")
    sw.input_cell(r + 1, 4, defaults.get("days_month", 21), num=FMT_NUM)
    sw.mark("days_month", r + 1)
    sw.label(r + 2, "Months per year (for yearly basis)")
    sw.input_cell(r + 2, 4, 12, num=FMT_NUM)
    sw.mark("months_yr", r + 2)
    sw.note(r, "Used only when rate basis = Hourly.", c=6)
    sw.skip(4)

    # --- 2. role table / blended rate ------------------------------------
    sw.section("2. STANDARD TEAM COMPOSITION & BLENDED RATE")
    sw.skip(1)
    hdr = sw.r
    for ci, h in [(2, "Role"), (3, "# in std team"), (4, "Yearly rate ($/yr)"),
                  (5, "Monthly rate ($/mo)"), (6, "Hourly rate ($/hr)"),
                  (7, "Monthly used ($/mo)"), (8, "Team cost ($/mo)"),
                  (9, "Source / notes")]:
        sw.cell(hdr, ci, h, f(bold=True, size=10, color=WHITE), fl=DARKBLUE_FILL,
                wrap=True, align="center")
    sw.skip(1)
    first = sw.r
    roles = defaults["roles"]
    n_spare = 3
    for i, (role, count, annual) in enumerate(roles + [("", 0, 0)] * n_spare):
        r = sw.r
        is_spare = i >= len(roles)
        sw.input_cell(r, 2, role if role else None)
        sw.input_cell(r, 3, count if not is_spare else None, num=FMT_NUM)
        sw.input_cell(r, 4, annual if not is_spare else None, num=FMT_USD)
        # monthly + hourly derived from yearly by default, but editable pattern:
        # keep D (yearly) as canonical input; E/F derived unless overridden later
        sw.formula_cell(r, 5, f"=IF(D{r}=\"\",0,D{r}/$D${sw.a('months_yr')})", num=FMT_USD)
        sw.formula_cell(
            r, 6,
            f"=IF(D{r}=\"\",0,E{r}/($D${sw.a('hours_day')}*$D${sw.a('days_month')}))",
            num=FMT_USD2)
        sw.formula_cell(
            r, 7,
            f"=IF($D${sw.a('basis')}=\"Hourly\",F{r}*$D${sw.a('hours_day')}*$D${sw.a('days_month')},"
            f"IF($D${sw.a('basis')}=\"Monthly\",E{r},"
            f"IF(D{r}=\"\",0,D{r}/$D${sw.a('months_yr')})))", num=FMT_USD)
        sw.formula_cell(r, 8, f"=IF(C{r}=\"\",0,C{r}*G{r})", num=FMT_USD)
        if is_spare:
            sw.cell(r, 9, "spare row — add roles as team shape is refined",
                    f(italic=True, size=9, color="FF808080"))
        sw.skip(1)
    last = sw.r - 1
    tr = sw.r
    sw.cell(tr, 2, "Total standard team", f(bold=True))
    sw.formula_cell(tr, 3, f"=SUM(C{first}:C{last})", num=FMT_NUM)
    sw.formula_cell(tr, 8, f"=SUM(H{first}:H{last})", num=FMT_USD)
    sw.mark("team_total", tr)
    sw.skip(2)
    r = sw.r
    sw.cell(r, 2, "BLENDED RATE — per team member", f(bold=True, size=12), fl="FFE2EFDA")
    sw.formula_cell(r, 7, f"=IF(C{tr}=0,0,H{tr}/C{tr})", num=FMT_USD)
    sw.cell(r, 8, "$/person/month", f(italic=True, size=9))
    sw.mark("blended_monthly", r)
    sw.formula_cell(r + 1, 7, f"=G{r}*$D${sw.a('months_yr')}", num=FMT_USD)
    sw.cell(r + 1, 8, "$/person/year", f(italic=True, size=9))
    sw.formula_cell(
        r + 2, 7, f"=G{r}/($D${sw.a('hours_day')}*$D${sw.a('days_month')})",
        num=FMT_USD2)
    sw.cell(r + 2, 8, "$/person/hour", f(italic=True, size=9))
    sw.note(r, "Weighted average monthly cost of one member of the standard team. "
               "Initiative milestone tables use this unless a rate override is entered.", c=9)
    sw.skip(4)

    # --- 3. contingency reference ----------------------------------------
    sw.section("3. CONTINGENCY FACTOR REFERENCE")
    sw.skip(1)
    r = sw.r
    sw.cell(r, 2, "Complexity / risk level", f(bold=True, size=10, color=WHITE), fl=DARKBLUE_FILL)
    sw.cell(r, 3, "Factor", f(bold=True, size=10, color=WHITE), fl=DARKBLUE_FILL, align="center")
    sw.cell(r, 6, "Based on the complexity and risks to deliver the milestone.",
            f(italic=True, size=9, color="FF595959"))
    sw.skip(1)
    sw.mark("cont_first", sw.r)
    for name, pct in CONT_TABLE:
        sw.label(sw.r, name)
        sw.input_cell(sw.r, 3, pct, num=FMT_PCT0)
        sw.skip(1)
    sw.mark("cont_last", sw.r - 1)
    sw.skip(2)

    # --- 4. capitalization policy ----------------------------------------
    sw.section("4. CAPITALIZATION POLICY (labor) — default % of milestone labor capitalizable")
    sw.skip(1)
    r = sw.r
    sw.cell(r, 2, "Milestone", f(bold=True, size=10, color=WHITE), fl=DARKBLUE_FILL)
    sw.cell(r, 3, "Cap %", f(bold=True, size=10, color=WHITE), fl=DARKBLUE_FILL, align="center")
    sw.cell(r, 6, "Share of labor cost treated as CapEx (feeds WAVE 'Capital Expenditures'); "
                  "remainder is expensed one-time cost (WAVE 'OTC'). Defaults follow a typical "
                  "internal-use software policy (ASC 350-40): application development/build stage "
                  "capitalizable, planning and post-launch expensed. CONFIRM WITH FINANCE — "
                  "initiative models reference these cells, override there if needed.",
            f(italic=True, size=9, color="FF595959"), wrap=True)
    ws.row_dimensions[r].height = 45
    sw.skip(1)
    cap_rows = {}
    for m in MILESTONES:
        sw.label(sw.r, m)
        sw.input_cell(sw.r, 3, defaults.get("cap_pcts", {}).get(m, 0.0), num=FMT_PCT0)
        cap_rows[m] = sw.r
        sw.skip(1)
    sw.label(sw.r, "Non-labor one-time (default)")
    sw.input_cell(sw.r, 3, defaults.get("cap_nonlabor", 0.0), num=FMT_PCT0)
    cap_rows["nonlabor"] = sw.r
    sw.skip(2)
    if defaults.get("notes"):
        sw.section("5. NOTES")
        for n in defaults["notes"]:
            sw.cell(sw.r, 2, "• " + n, f(size=10), wrap=True)
            sw.skip(1)

    anchors = {k: sw.a(k) for k in ["basis", "hours_day", "days_month",
                                     "months_yr", "blended_monthly",
                                     "team_total"]}
    anchors["cap_rows"] = cap_rows
    anchors["sheet"] = ws_name
    return anchors


# ------------------------------------------------------------------- Cover
def build_cover(wb, meta):
    ws = wb.create_sheet("Cover", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "1F3864"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 100
    ws.column_dimensions["D"].width = 30
    sw = SheetWriter(ws)
    sw.r = 2
    sw.cell(2, 2, meta["title"], f(bold=True, size=20, color=WHITE), fl=DARKBLUE_FILL)
    sw.cell(2, 3, None, fl=DARKBLUE_FILL)
    sw.cell(2, 4, None, fl=DARKBLUE_FILL)
    ws.row_dimensions[2].height = 30
    sw.r = 4
    for k, v in [("Workstream", meta["workstream"]),
                 ("Purpose", meta["purpose"]),
                 ("Version / status", meta["version"]),
                 ("Prepared", meta["prepared"]),
                 ("Sources", meta["sources"])]:
        sw.label(sw.r, k, bold=True)
        sw.cell(sw.r, 3, v, f(), wrap=True)
        sw.skip(1)
    sw.skip(1)
    sw.section("HOW THIS FILE IS ORGANIZED")
    sw.skip(1)
    for name, desc in meta["sheet_index"]:
        sw.cell(sw.r, 2, name, f(bold=True, color="FF1F3864"))
        sw.cell(sw.r, 3, desc, f(size=10), wrap=True)
        sw.skip(1)
    sw.skip(1)
    sw.section("COLOR LEGEND")
    sw.skip(1)
    legend = [
        ("Input — change freely", f(color=BLUE), INPUT_FILL),
        ("Formula — do not overwrite", f(), None),
        ("Link to another sheet", f(color=GREEN), None),
        ("Section band", f(bold=True, color="FF1F3864"), BAND_FILL),
    ]
    for text, fnt, flc in legend:
        sw.cell(sw.r, 2, text, fnt, fl=flc, border=BORDER_ALL if flc else None)
        sw.skip(1)
    sw.skip(1)
    sw.section("MODEL -> WAVE LINKING")
    sw.skip(1)
    for t in meta["linking_notes"]:
        sw.cell(sw.r, 2, "• " + t, f(size=10), wrap=True)
        ws.row_dimensions[sw.r].height = 26
        sw.skip(1)
    sw.skip(1)
    sw.section("DEVIATIONS FROM THE STANDARD WAVE TEMPLATE (documented fixes)")
    sw.skip(1)
    for t in meta["deviations"]:
        sw.cell(sw.r, 2, "• " + t, f(size=10), wrap=True)
        ws.row_dimensions[sw.r].height = 26
        sw.skip(1)
    sw.skip(1)
    sw.section("OPEN ITEMS / QC FLAGS CARRIED FROM SOURCE MATERIALS")
    sw.skip(1)
    for t in meta["open_items"]:
        sw.cell(sw.r, 2, "• " + t, f(size=10), wrap=True)
        ws.row_dimensions[sw.r].height = 26
        sw.skip(1)
    return ws
