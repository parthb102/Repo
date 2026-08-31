#!/usr/bin/env python3
"""Model sheet engine + WAVE filler + Summary builder."""
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from framework import (SheetWriter, copy_wave_sheet, wave_write, f, fill, col,
                       FONT, BLUE, BLACK, GREEN, GREY_FILL, DARKBLUE_FILL,
                       BAND_FILL, INPUT_FILL, WHITE, FMT_USD, FMT_USD2,
                       FMT_USDM, FMT_PCT, FMT_PCT0, FMT_NUM, FMT_NUM1,
                       FMT_DATE, MON_FIRST, MON_LAST, YR_FIRST, YEARS,
                       BORDER_ALL)
from builders import MILESTONES, _model_cols

Q_BU_COL = 17  # Q


def _date_formula(ym):
    y, m = ym
    return f"=DATE({y},{m},1)"


class ModelBuilder:
    """Builds one '<code> Model' sheet from a spec; records feed row map."""

    def __init__(self, wb, spec, rc):
        self.wb = wb
        self.spec = spec
        self.rc = rc  # rate card anchors
        self.ws = wb.create_sheet(spec["model_sheet"])
        self.ws.sheet_view.showGridLines = False
        self.ws.sheet_properties.tabColor = spec.get("tab_color", "4472C4")
        _model_cols(self.ws)
        self.ws.column_dimensions["C"].width = 34
        self.sw = SheetWriter(self.ws)
        self.sym = {}       # symbol -> absolute ref (same-sheet)
        self.feed = {"benefit": [], "runcost": [], "capex": [], "otc": [],
                     "model_sheet": spec["model_sheet"]}

    # -- helpers ----------------------------------------------------------
    def ref(self, sym):
        return self.sym[sym]

    def sub(self, formula):
        out = formula
        for k, v in self.sym.items():
            out = out.replace("{" + k + "}", v)
        out = out.replace("{rc}", f"'{self.rc['sheet']}'")
        out = out.replace("{rc_blended}",
                          f"'{self.rc['sheet']}'!$G${self.rc['blended_monthly']}")
        return out

    def _grid_monthly(self, r, formula_tpl, num=FMT_USDM):
        """Write monthly T:CA cells; formula_tpl may contain {c} (col letter),
        {yr} (year helper cell), {hd} (date header cell)."""
        for i in range(60):
            c = MON_FIRST + i
            cl = col(c)
            self.sw.formula_cell(
                r, c,
                formula_tpl.format(c=cl, yr=f"{cl}${self.yr_row}",
                                   hd=f"{cl}${self.hd_row}"),
                num=num)

    def _grid_annual(self, r, num=FMT_USDM, with_k=True):
        for i, y in enumerate(YEARS):
            c = YR_FIRST + i
            self.sw.formula_cell(
                r, c, f"=SUMIFS($T{r}:$CA{r},$T${self.yr_row}:$CA${self.yr_row},{y})",
                num=num)
        if with_k:
            self.sw.formula_cell(
                r, YR_FIRST - 1,
                f"=SUMIFS($T{r}:$CA{r},$T${self.yr_row}:$CA${self.yr_row},2030)",
                num=num, comment="Annual run rate = 2030 value (matches WAVE check rule)")

    # -- sections ---------------------------------------------------------
    def header(self):
        sp, sw = self.spec, self.sw
        sw.r = 2
        sw.title_row(f"{sp['code']} — {sp['short']}  |  INITIATIVE MODEL")
        sw.cell(sw.r, 2,
                f"Workstream: {sp['workstream']}   •   Model status: {sp['status']}",
                f(italic=True, size=10, color="FF595959"))
        sw.skip(2)
        # month header pair
        self.yr_row = sw.r
        self.hd_row = sw.r + 1
        from framework import month_header_rows
        month_header_rows(sw, self.yr_row, self.hd_row)
        sw.cell(self.hd_row, 2, "Monthly grid (Jan-26 → Dec-30) →",
                f(italic=True, size=9, color="FF808080"))
        sw.r = self.hd_row + 1
        self.ws.freeze_panes = f"G{self.hd_row + 1}"
        sw.skip(1)

    def status_block(self):
        sp, sw = self.spec, self.sw
        sw.section("0. INITIATIVE & STATUS")
        sw.skip(1)
        r = sw.r
        rows = [("Initiative name", sp["name"], None),
                ("Initiative owner", sp["owner"], "confirm with workstream lead"),
                ("Impacted BU", sp["bu"], None)]
        for lab, val, note in rows:
            sw.label(sw.r, lab)
            sw.input_cell(sw.r, 4, val)
            self.ws.merge_cells(start_row=sw.r, start_column=4, end_row=sw.r, end_column=8)
            if note:
                sw.note(sw.r, note, c=9)
            sw.skip(1)
        sw.label(sw.r, "Go-live / Estimated L4 date", bold=True)
        sw.input_cell(sw.r, 4, _date_formula(sp["golive"]), num=FMT_DATE,
                      comment="Benefits start in this month; feeds the WAVE 'Estimated L4 date'")
        self.sym["golive"] = f"$D${sw.r}"
        sw.note(sw.r, "Benefit ramp is gated on this date.", c=9)
        sw.skip(2)
        for s in sp.get("sources", []):
            sw.cell(sw.r, 2, "Source: " + s, f(italic=True, size=9, color="FF595959"))
            sw.skip(1)
        if sp.get("open_items"):
            sw.cell(sw.r, 2, "Open items to refine with initiative owner:",
                    f(bold=True, size=9, color="FFC00000"))
            sw.skip(1)
            for s in sp["open_items"]:
                sw.cell(sw.r, 2, "  – " + s, f(size=9, color="FFC00000"))
                sw.skip(1)
        sw.skip(1)

    def _table_rows(self, rows, is_input):
        """Write assumption/calc rows. Each row: (sym,label,value,unit,fmt,note)
        or ('#SUB', text)."""
        sw = self.sw
        for row in rows:
            if row[0] == "#SUB":
                sw.cell(sw.r, 2, row[1], f(bold=True, italic=True, size=10,
                                           color="FF1F3864"))
                sw.skip(1)
                continue
            sym, label, value, unit, fmt, note = row
            sw.label(sw.r, label, indent=1)
            if value is not None:
                if is_input and not (isinstance(value, str) and value.startswith("=")):
                    sw.input_cell(sw.r, 4, value, num=fmt)
                elif is_input:
                    sw.input_cell(sw.r, 4, self.sub(value), num=fmt)
                else:
                    v = self.sub(value) if isinstance(value, str) else value
                    link = isinstance(v, str) and "'" in v
                    sw.formula_cell(sw.r, 4, v, num=fmt, link=link)
            if unit:
                sw.cell(sw.r, 5, unit, f(italic=True, size=9, color="FF595959"))
            if note:
                sw.note(sw.r, note, c=6)
            if sym:
                self.sym[sym] = f"$D${sw.r}"
            sw.skip(1)

    def assumptions(self):
        sw = self.sw
        sw.section("1. SIZING ASSUMPTIONS  (blue = inputs — refine with initiative owner)")
        sw.skip(1)
        self._table_rows(self.spec["assumptions"], is_input=True)
        for i in range(2):
            sw.label(sw.r, "(spare assumption row)", indent=1)
            self.sw.input_cell(sw.r, 4, None)
            sw.skip(1)
        sw.skip(1)

    def calc(self):
        sw = self.sw
        sw.section("2. BENEFIT CALCULATION  (steady-state annual run-rate)")
        sw.skip(1)
        self._table_rows(self.spec["calc"], is_input=False)
        sw.skip(1)

    def ramp(self):
        sw, sp = self.sw, self.spec
        sw.section("3. RAMP & PHASING")
        sw.skip(1)
        r = sw.r
        sw.label(r, "Months to full run rate after go-live (linear ramp)", bold=True)
        sw.input_cell(r, 4, sp.get("ramp_months", 6), num=FMT_NUM,
                      comment="Benefit ramps linearly from 0% at go-live to 100% after this many months "
                              "(same convention as the Servicing S2 model: scale-up phase drives the ramp).")
        self.sym["ramp_months"] = f"$D${r}"
        sw.note(r, "Benefit starts at go-live and ramps linearly to steady state over this many months.", c=6)
        sw.skip(1)
        r = sw.r
        sw.label(r, "Monthly ramp % (run-rate captured in month)")
        g = self.sym["golive"]
        rm = self.sym["ramp_months"]
        self._grid_monthly(
            r,
            "=MIN(1,MAX(0,((YEAR({hd})-YEAR(" + g + "))*12+MONTH({hd})-MONTH(" + g + ")+0.5)/"
            "MAX(" + rm + ",0.5)))",
            num=FMT_PCT0)
        self.ramp_month_row = r
        sw.skip(1)
        r = sw.r
        sw.label(r, "Resulting average ramp % by year (info)")
        for i, y in enumerate(YEARS):
            c = YR_FIRST + i
            sw.formula_cell(
                r, c,
                f"=SUMIFS($T{self.ramp_month_row}:$CA{self.ramp_month_row},"
                f"$T${self.yr_row}:$CA${self.yr_row},{y})/12", num=FMT_PCT0)
        sw.skip(2)

    def benefits(self):
        sw, sp = self.sw, self.spec
        sw.section("4. RECURRING BENEFIT LINES  ($M — these rows feed the WAVE sheet)")
        sw.skip(1)
        hr = sw.r
        sw.cell(hr, 2, "Line (WAVE category)", f(bold=True, size=9, color=WHITE), fl=DARKBLUE_FILL)
        sw.cell(hr, 4, "Baseline $M", f(bold=True, size=9, color=WHITE), fl=DARKBLUE_FILL, align="center", wrap=True)
        sw.cell(hr, 11, "Run rate", f(bold=True, size=9, color=WHITE), fl=DARKBLUE_FILL, align="center")
        sw.skip(1)
        first = sw.r
        for line in sp["benefit_lines"]:
            r = sw.r
            sw.label(r, line["label"] + f"   [{line['category'].upper()}]", bold=False)
            base = line.get("baseline")
            if base is not None:
                if isinstance(base, str):
                    sw.formula_cell(r, 4, self.sub(base), num=FMT_USDM)
                else:
                    sw.input_cell(r, 4, base, num=FMT_USDM)
            else:
                sw.input_cell(r, 4, 0, num=FMT_USDM,
                              comment="Baseline annual spend/revenue of this line, $M (for WAVE col I)")
            self._grid_monthly(
                r,
                "=(" + self.sub(line["runrate"]) + ")/1000000/12*{c}$" + str(self.ramp_month_row))
            self._grid_annual(r)
            if line.get("comment"):
                sw.cell(r, 18, None)
            self.feed["benefit"].append(
                dict(row=r, label=line["label"], category=line["category"],
                     comment=line.get("comment", ""), sym=line.get("sym")))
            sw.skip(1)
        last = sw.r - 1
        r = sw.r
        sw.cell(r, 2, "Total recurring benefit ($M)", f(bold=True))
        for c in range(YR_FIRST - 1, YR_FIRST + 5):
            sw.formula_cell(r, c, f"=SUM({col(c)}{first}:{col(c)}{last})", num=FMT_USDM)
        self.feed["benefit_total_row"] = r
        sw.skip(2)

    def investment(self):
        sw, sp = self.sw, self.spec
        sw.section("5. ONE-TIME INVESTMENT — people cost by milestone (standard approach)")
        sw.skip(1)
        blocks = sp.get("milestone_blocks") or [
            {"title": None, "milestones": sp.get("milestones", {})}]
        cap_rows = self.rc["cap_rows"]
        block_ranges = []   # (m_first, m_last) per block
        for bi, block in enumerate(blocks):
            if block.get("title"):
                sw.cell(sw.r, 2, block["title"], f(bold=True, size=10, color="FF1F3864"))
                sw.skip(1)
            hr = sw.r
            heads = [(2, "Milestone"), (3, "Assumptions"), (4, "Team size"),
                     (5, "Billing months"), (6, "Contingency Factor"),
                     (7, "Labor Costs"), (8, "Start month"), (9, "Cap % (labor)"),
                     (11, "Rate override ($/mo)")]
            for c, h in heads:
                sw.cell(hr, c, h, f(bold=True, size=9, color=WHITE), fl=DARKBLUE_FILL,
                        wrap=True, align="center")
            sw.skip(1)
            if bi == 0:
                sub = sw.r
                sw.cell(sub, 3, "Capture any assumptions on cost drivers, timelines, contingency factor etc.",
                        f(italic=True, size=8, color="FF808080"), wrap=True)
                sw.cell(sub, 5, "Months the team works at full utilization",
                        f(italic=True, size=8, color="FF808080"), wrap=True)
                sw.cell(sub, 6, "Low=10% Medium=15% High=20% Very high=50%",
                        f(italic=True, size=8, color="FF808080"), wrap=True)
                sw.cell(sub, 7, "Total labor costs for milestone",
                        f(italic=True, size=8, color="FF808080"), wrap=True)
                sw.cell(sub, 9, "Default links to Rate Card policy — type a % to override",
                        f(italic=True, size=8, color="FF808080"), wrap=True)
                self.ws.row_dimensions[sub].height = 30
                sw.skip(1)
            m_first = sw.r
            for m in MILESTONES:
                ms = block["milestones"].get(m, {})
                r = sw.r
                sw.cell(r, 2, m, f(bold=True), border=BORDER_ALL)
                sw.input_cell(r, 3, ms.get("assump", ""), )
                self.ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
                sw.input_cell(r, 4, ms.get("team", 0), num=FMT_NUM1)
                sw.input_cell(r, 5, ms.get("months", 0), num=FMT_NUM1)
                sw.input_cell(r, 6, ms.get("cont", 0.10), num=FMT_PCT0,
                              comment="Low=10% Medium=15% High=20% Very high=50% (see Rate Card)")
                rate_override = ms.get("rate_override")
                sw.formula_cell(r, 7,
                                f"=D{r}*E{r}*IF(K{r}=\"\",{{rc_blended}},K{r})*(1+F{r})".replace(
                                    "{rc_blended}",
                                    f"'{self.rc['sheet']}'!$G${self.rc['blended_monthly']}"),
                                num=FMT_USD)
                sw.input_cell(r, 8, _date_formula(ms.get("start", sp["golive"])), num=FMT_DATE)
                sw.formula_cell(r, 9,
                                f"='{self.rc['sheet']}'!$C${cap_rows[m]}",
                                num=FMT_PCT0, link=True,
                                comment="Default = Rate Card capitalization policy. Type a % to override.")
                if rate_override is not None:
                    sw.input_cell(r, 11, rate_override, num=FMT_USD)
                else:
                    sw.input_cell(r, 11, None, num=FMT_USD)
                self.ws.row_dimensions[r].height = 42
                sw.skip(1)
            m_last = sw.r - 1
            block_ranges.append((m_first, m_last))
            sw.skip(1)

        # monthly phasing rows ($) for every block
        sw.cell(sw.r, 2, "Milestone labor phasing ($/month, spread evenly over billing months)",
                f(italic=True, size=9, color="FF595959"))
        sw.skip(1)
        phase_ranges = []
        for bi, (m_first, m_last) in enumerate(block_ranges):
            p_first = sw.r
            for i, m in enumerate(MILESTONES):
                r = sw.r
                mr = m_first + i
                lab = m if len(blocks) == 1 else f"{m}  ({blocks[bi].get('title') or 'block ' + str(bi+1)})"
                sw.cell(r, 2, "   " + lab, f(size=9))
                self._grid_monthly(
                    r,
                    "=IF(AND({hd}>=EOMONTH($H$" + str(mr) + ",0),"
                    "{hd}<=EOMONTH(EDATE($H$" + str(mr) + ",$E$" + str(mr) + "-1),0)),"
                    "$G$" + str(mr) + "/MAX($E$" + str(mr) + ",1),0)",
                    num=FMT_USD)
                sw.skip(1)
            phase_ranges.append((p_first, sw.r - 1))
        # grand total labor
        r = sw.r
        sw.cell(r, 2, "Total milestone labor (all blocks)", f(bold=True))
        parts = "+".join(f"SUM(G{a}:G{b})" for a, b in block_ranges)
        sw.formula_cell(r, 7, f"={parts}", num=FMT_USD)
        self._ms_total_row = r
        self._block_ranges = block_ranges
        self._phase_ranges = phase_ranges
        sw.skip(2)

        # non-labor table
        sw.cell(sw.r, 2, "One-time non-labor costs (vendor, integration, licences, change mgmt., etc.)",
                f(bold=True, size=10))
        sw.skip(1)
        hr = sw.r
        for c, h in [(2, "Item"), (3, "Assumptions / notes"), (4, "Total $"),
                     (5, "Spread months"), (8, "Start month"), (9, "Cap %")]:
            sw.cell(hr, c, h, f(bold=True, size=9, color=WHITE), fl=DARKBLUE_FILL,
                    wrap=True, align="center")
        sw.skip(1)
        nl_first = sw.r
        items = sp.get("onetime_nonlabor", [])
        for k, it in enumerate(items + [dict(label="(spare)", total=0, months=1,
                                             start=sp["golive"], cap=0.0, notes="")]):
            r = sw.r
            sw.input_cell(r, 2, it["label"])
            sw.input_cell(r, 3, it.get("notes", ""))
            self.ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
            tot = it.get("total", 0)
            if isinstance(tot, str):
                ftxt = self.sub(tot).replace("{MS_TOTAL}", f"$G${self._ms_total_row}")
                for j in range(k):
                    ftxt = ftxt.replace("{NL:%d}" % (j + 1), f"$D${nl_first + j}")
                sw.formula_cell(r, 4, ftxt, num=FMT_USD)
            else:
                sw.input_cell(r, 4, tot, num=FMT_USD)
            sw.input_cell(r, 5, it.get("months", 1), num=FMT_NUM)
            sw.input_cell(r, 8, _date_formula(it.get("start", sp["golive"])), num=FMT_DATE)
            sw.input_cell(r, 9, it.get("cap", 0.0), num=FMT_PCT0)
            sw.skip(1)
        nl_last = sw.r - 1
        sw.skip(1)
        sw.cell(sw.r, 2, "Non-labor phasing ($/month)", f(italic=True, size=9, color="FF595959"))
        sw.skip(1)
        q_first = sw.r
        for i in range(nl_last - nl_first + 1):
            r = sw.r
            nr = nl_first + i
            sw.cell(r, 2, f"   phasing — item {i+1}", f(size=9))
            self._grid_monthly(
                r,
                "=IF(AND({hd}>=EOMONTH($H$" + str(nr) + ",0),"
                "{hd}<=EOMONTH(EDATE($H$" + str(nr) + ",$E$" + str(nr) + "-1),0)),"
                "$D$" + str(nr) + "/MAX($E$" + str(nr) + ",1),0)",
                num=FMT_USD)
            sw.skip(1)
        q_last = sw.r - 1
        self._nl_rows = (nl_first, nl_last, q_first, q_last)
        sw.skip(1)

        # feed rows ($M)
        sw.cell(sw.r, 2, "One-time investment feed rows ($M — these feed the WAVE sheet)",
                f(bold=True, size=10))
        sw.skip(1)
        nl_first, nl_last, q_first, q_last = self._nl_rows
        lab_cap = "+".join(
            f"SUMPRODUCT({{c}}{p1}:{{c}}{p2},$I${m1}:$I${m2})"
            for (m1, m2), (p1, p2) in zip(self._block_ranges, self._phase_ranges))
        lab_exp = "+".join(
            f"SUMPRODUCT({{c}}{p1}:{{c}}{p2},1-$I${m1}:$I${m2})"
            for (m1, m2), (p1, p2) in zip(self._block_ranges, self._phase_ranges))
        specs = [
            ("CAPEX — capitalized build labor", "capex",
             f"=({lab_cap})/1000000"),
            ("CAPEX — capitalized non-labor", "capex",
             f"=SUMPRODUCT({{c}}{q_first}:{{c}}{q_last},$I${nl_first}:$I${nl_last})/1000000"),
            ("OTC — build labor (expensed)", "otc",
             f"=({lab_exp})/1000000"),
            ("OTC — non-labor (expensed)", "otc",
             f"=SUMPRODUCT({{c}}{q_first}:{{c}}{q_last},1-$I${nl_first}:$I${nl_last})/1000000"),
        ]
        for label, kind, tpl in specs:
            r = sw.r
            sw.label(r, label)
            self._grid_monthly(r, tpl)
            self._grid_annual(r, with_k=False)
            self.feed[kind].append(dict(row=r, label=label))
            sw.skip(1)
        r = sw.r
        sw.cell(r, 2, "Total one-time investment ($M)", f(bold=True))
        for c in range(YR_FIRST, YR_FIRST + 5):
            sw.formula_cell(r, c, f"=SUM({col(c)}{r-4}:{col(c)}{r-1})", num=FMT_USDM)
        self.feed["onetime_total_row"] = r
        sw.skip(2)

    def runcosts(self):
        sw, sp = self.sw, self.spec
        sw.section("6. RECURRING RUN COSTS to achieve impact")
        sw.skip(1)
        hr = sw.r
        for c, h in [(2, "Line"), (3, "Assumptions / notes"), (4, "Annual $ (steady state)"),
                     (8, "Start month"), (9, "Scales with ramp? (Y/N)")]:
            sw.cell(hr, c, h, f(bold=True, size=9, color=WHITE), fl=DARKBLUE_FILL,
                    wrap=True, align="center")
        sw.skip(1)
        rc_first = sw.r
        lines = sp.get("runcost_lines", [])
        if lines:
            from openpyxl.worksheet.datavalidation import DataValidation
            dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
            self.ws.add_data_validation(dv)
            dv.add(f"I{rc_first}:I{rc_first + len(lines) - 1}")
        for it in lines:
            r = sw.r
            sw.label(r, it["label"])
            sw.input_cell(r, 3, it.get("notes", ""))
            self.ws.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
            ann = it.get("annual", 0)
            if isinstance(ann, str):
                sw.formula_cell(r, 4, self.sub(ann), num=FMT_USD)
            else:
                sw.input_cell(r, 4, ann, num=FMT_USD)
            sw.input_cell(r, 8, _date_formula(it.get("start", sp["golive"])), num=FMT_DATE)
            sw.input_cell(r, 9, it.get("scale", "N"))
            sw.skip(1)
        rc_last = sw.r - 1
        sw.skip(1)
        sw.cell(sw.r, 2, "Run-cost feed rows ($M)", f(bold=True, size=10))
        sw.skip(1)
        for i, it in enumerate(lines):
            r = sw.r
            nr = rc_first + i
            sw.label(r, it["label"] + "   [RUN]")
            self._grid_monthly(
                r,
                "=$D$" + str(nr) + "/1000000/12*IF($I$" + str(nr) + "=\"Y\","
                "{c}$" + str(self.ramp_month_row) + ","
                "IF({hd}>=EOMONTH($H$" + str(nr) + ",0),1,0))")
            self._grid_annual(r)
            self.feed["runcost"].append(dict(row=r, label=it["label"],
                                             comment=it.get("notes", "")))
            sw.skip(1)
        r = sw.r
        sw.cell(r, 2, "Total run cost ($M)", f(bold=True))
        n = len(lines)
        if n:
            for c in range(YR_FIRST - 1, YR_FIRST + 5):
                sw.formula_cell(r, c, f"=SUM({col(c)}{r-n}:{col(c)}{r-1})", num=FMT_USDM)
        self.feed["runcost_total_row"] = r
        sw.skip(2)

    def mapping_note(self):
        sw = self.sw
        sw.section("7. WAVE MAPPING (reference)")
        sw.skip(1)
        rows = []
        cat_slots = {"rev": [66, 67], "cos": [71, 72, 73, 74, 75],
                     "opex": [79, 80, 81, 82, 83]}
        slots = {k: list(v) for k, v in cat_slots.items()}
        for b in self.feed["benefit"]:
            tgt = slots[b["category"]].pop(0)
            b["wave_row"] = tgt
            rows.append((b["row"], b["label"], f"WAVE row {tgt} (Section II — Recurring Benefits)"))
        run_slots = [107, 108, 109, 110, 111]
        for b in self.feed["runcost"]:
            tgt = run_slots.pop(0)
            b["wave_row"] = tgt
            rows.append((b["row"], b["label"], f"WAVE row {tgt} (Section II — Recurring Costs)"))
        for b, tgt in zip(self.feed["capex"], [136, 137]):
            b["wave_row"] = tgt
            rows.append((b["row"], b["label"], f"WAVE row {tgt} (Section III — CapEx)"))
        for b, tgt in zip(self.feed["otc"], [141, 142]):
            b["wave_row"] = tgt
            rows.append((b["row"], b["label"], f"WAVE row {tgt} (Section III — OTC)"))
        for mr, lab, tgt in rows:
            sw.cell(sw.r, 2, f"Model row {mr} — {lab}", f(size=9))
            sw.cell(sw.r, 6, "→  " + tgt, f(size=9, color="FF1F3864"))
            sw.skip(1)
        sw.skip(1)
        sw.section("8. FLEX / SCRATCH AREA — free space for refinements with the initiative owner")
        sw.skip(6)

    def build(self):
        self.header()
        self.status_block()
        self.assumptions()
        if self.spec.get("custom_section"):
            self.spec["custom_section"](self)
        self.calc()
        self.ramp()
        self.benefits()
        self.investment()
        self.runcosts()
        self.mapping_note()
        return self.feed


# ------------------------------------------------------------------ WAVE
def fill_wave(wb, spec, feed, sym):
    ws = copy_wave_sheet(wb, spec["wave_sheet"])
    m = f"'{spec['model_sheet']}'"

    # template bug fixes (documented on Cover)
    ws["K88"] = "=SUM(K68, K76, K84)"
    ws["G163"] = '=IFERROR(K117/-SUM(L149:P149),"n/a")'
    ws["G164"] = '=IFERROR(IRR($L$159:$P$159, 0.2),"n/a")'
    ws["I64"] = "=$H$18"
    ws["I92"] = "=$H$18"
    # payback chain: template stalls when a year's cumulative CF is exactly 0
    ws["L161"] = "=IF(L160<0,1,0)"
    for cur, prev in [("M", "L"), ("N", "M"), ("O", "N"), ("P", "O")]:
        ws[f"{cur}161"] = (f"=IF({cur}160<0,1,IF(AND({prev}160<0,{cur}159>0),"
                           f"-{prev}160/{cur}159,0))")
    ws["G165"] = '=IF(SUM(L161:P161)=5,">5",SUM(L161:P161))'
    # neutralize template illustrative prefills
    ws["I79"] = 0
    ws["J107"] = 0

    # header inputs
    wave_write(ws, "F6", spec.get("wave_number", "TBD"))
    wave_write(ws, "F7", spec["owner"])
    wave_write(ws, "F8", spec["name"])
    ws["F11"] = f"={m}!{sym['golive']}"
    ws["F11"].number_format = "m/d/yyyy"
    _greenify(ws, "F11")

    # Section I — driver summary
    drv = spec.get("wave_driver")
    ws["D16"] = ("Primary value driver — linked to the Model sheet; refine with "
                 "the initiative owner")
    if drv:
        base_rows = [18, 19, 20, 21]
        goal_rows = [24, 25, 26]
        incr_rows = [29]
        _fill_driver_block(ws, m, drv.get("baseline", []), base_rows, sym)
        _fill_driver_block(ws, m, drv.get("goal", []), goal_rows, sym)
        _fill_driver_block(ws, m, drv.get("incr", []), incr_rows, sym)
    # clear placeholder blocks 2 & 3 (rows 31..57)
    for r in range(31, 58):
        for c in ["D", "E", "G", "H", "I", "J"]:
            cc = ws[f"{c}{r}"]
            cc.value = None

    # Section II — benefits
    # K is linked to the model's run-rate cell directly (not recomputed as
    # I-J / J-I): J is derived from I and that same cell, so the template's
    # arithmetic still holds, while K and P share the identical double and
    # the R-column run-rate checks stay exact in Excel (no last-bit float
    # differences from subtract-and-restore).
    for b in feed["benefit"]:
        r = b["wave_row"]
        mr = b["row"]
        ws[f"E{r}"] = b["label"]
        if b.get("comment"):
            ws[f"H{r}"] = b["comment"]
        ws[f"I{r}"] = f"={m}!$D${mr}"
        _greenify(ws, f"I{r}")
        if b["category"] == "rev":
            ws[f"J{r}"] = f"=I{r}+{m}!$K${mr}"
        else:
            ws[f"J{r}"] = f"=I{r}-{m}!$K${mr}"
        _greenify(ws, f"J{r}")
        ws[f"K{r}"] = f"={m}!$K${mr}"
        _greenify(ws, f"K{r}")
        _link_year_month(ws, m, r, mr)
        ws[f"Q{r}"] = spec["bu"]

    # Section II — recurring run costs
    for b in feed["runcost"]:
        r = b["wave_row"]
        mr = b["row"]
        ws[f"E{r}"] = b["label"]
        if b.get("comment"):
            ws[f"H{r}"] = b["comment"][:120]
        ws[f"I{r}"] = 0
        ws[f"J{r}"] = f"=I{r}+{m}!$K${mr}"
        _greenify(ws, f"J{r}")
        ws[f"K{r}"] = f"={m}!$K${mr}"
        _greenify(ws, f"K{r}")
        _link_year_month(ws, m, r, mr)
        ws[f"Q{r}"] = spec["bu"]

    # Section III — one-time
    for b in feed["capex"] + feed["otc"]:
        r = b["wave_row"]
        mr = b["row"]
        ws[f"E{r}"] = b["label"]
        _link_year_month(ws, m, r, mr)
        ws[f"Q{r}"] = spec["bu"]

    return ws


def _fill_driver_block(ws, m, rows, targets, sym):
    for i, tr in enumerate(targets):
        if i < len(rows):
            row = list(rows[i]) + [None] * 5
            label, val, unit, period, fmt = row[:5]
            ws[f"E{tr}"] = label
            if period is not None:
                ws[f"H{tr}"] = period
            ws[f"I{tr}"] = unit
            cc = ws[f"J{tr}"]
            if isinstance(val, str):
                v = val.replace("{m}", m)
                for name, ref in sym.items():
                    v = v.replace("{S:" + name + "}", f"{m}!{ref}")
                cc.value = v
            else:
                cc.value = val
            if fmt:
                cc.number_format = fmt
            if isinstance(val, str) and val.startswith("="):
                _greenify(ws, f"J{tr}")
            else:
                _blueify(ws, f"J{tr}")
        else:
            for c in ["E", "G", "H", "I", "J"]:
                ws[f"{c}{tr}"].value = None


def _link_year_month(ws, m, r, mr):
    from framework import YR_FIRST, MON_FIRST
    for i in range(5):
        c = get_column_letter(YR_FIRST + i)
        ws[f"{c}{r}"] = f"={m}!{c}{mr}"
        _greenify(ws, f"{c}{r}")
    for i in range(60):
        c = get_column_letter(MON_FIRST + i)
        ws[f"{c}{r}"] = f"={m}!{c}{mr}"
        _greenify(ws, f"{c}{r}")


def _greenify(ws, ref):
    from copy import copy as _copy
    cc = ws[ref]
    fnt = _copy(cc.font)
    fnt.color = "FF008000"
    cc.font = fnt


def _blueify(ws, ref):
    from copy import copy as _copy
    cc = ws[ref]
    fnt = _copy(cc.font)
    fnt.color = "FF0000FF"
    cc.font = fnt


# ---------------------------------------------------------------- Summary
def build_summary(wb, specs, title, cap_note=None):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "FFC000"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 44
    for c in "CDEFGHIJKL":
        ws.column_dimensions[c].width = 13
    sw = SheetWriter(ws)
    sw.r = 2
    sw.title_row(title + " — PORTFOLIO SUMMARY ($M)")
    sw.cell(sw.r, 2, "All figures link to the initiative WAVE sheets. Payback/NPV per WAVE Section IV.",
            f(italic=True, size=9, color="FF595959"))
    sw.skip(2)
    hr = sw.r
    heads = ["Initiative", "NRB run rate", "NRB 2026", "NRB 2027", "NRB 2028",
             "NRB 2029", "NRB 2030", "One-time cost (5yr)", "Run cost 2030",
             "NPV (5yr)", "Payback (yrs)"]
    for i, h in enumerate(heads):
        sw.cell(hr, 2 + i, h, f(bold=True, size=9, color=WHITE), fl=DARKBLUE_FILL,
                wrap=True, align="center")
    sw.skip(1)
    first = sw.r
    for sp in specs:
        r = sw.r
        w = f"'{sp['wave_sheet']}'"
        sw.cell(r, 2, f"{sp['code']} — {sp['short']}", f(bold=True))
        sw.formula_cell(r, 3, f"={w}!$K$117", num=FMT_USDM, link=True)
        for i, ycol in enumerate(["L", "M", "N", "O", "P"]):
            sw.formula_cell(r, 4 + i, f"={w}!${ycol}$117", num=FMT_USDM, link=True)
        sw.formula_cell(r, 9, f"=SUM({w}!$L$147:$P$147)", num=FMT_USDM, link=True)
        sw.formula_cell(r, 10, f"={w}!$P$115", num=FMT_USDM, link=True)
        sw.formula_cell(r, 11, f"={w}!$G$168", num=FMT_USDM, link=True)
        sw.formula_cell(r, 12, f"={w}!$G$165", num=FMT_NUM1, link=True)
        sw.skip(1)
    last = sw.r - 1
    r = sw.r
    sw.cell(r, 2, "TOTAL — " + title, f(bold=True), fl=BAND_FILL)
    for ci in range(3, 11):
        c = col(ci)
        sw.formula_cell(r, ci, f"=SUM({c}{first}:{c}{last})", num=FMT_USDM)
        ws.cell(row=r, column=ci).fill = fill(BAND_FILL)
        ws.cell(row=r, column=ci).font = f(bold=True)
    sw.skip(2)
    sw.cell(sw.r, 2, "Note: NPV and payback are not additive across initiatives; totals shown for benefit/cost columns only.",
            f(italic=True, size=9, color="FF595959"))
    if cap_note:
        sw.skip(1)
        sw.cell(sw.r, 2, cap_note, f(italic=True, size=9, color="FFC00000"))
    return ws
